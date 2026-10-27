import sys,json
from openpyxl import load_workbook
wb=load_workbook(sys.argv[1]);ws=wb.active;ws.delete_rows(2,ws.max_row)
for row in json.loads(sys.argv[3]):ws.append(row)
wb.save(sys.argv[2])
