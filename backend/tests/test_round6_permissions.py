"""第六轮真实 HTTP 权限矩阵、共用读取、迁移与隔离回归。"""
import asyncio, ast
from contextlib import asynccontextmanager
from datetime import datetime
from pathlib import Path
import pytest
from sqlalchemy import select
from app.core.menu_registry import MENU_PERMISSIONS_BY_KEY
from app.models import User, SystemSetting
from app.models.recharge import IncomingFlow, ReconciliationResult
from app.services import permission_migrations as migrations
from app.services.schedule import default_work_schedule
from tests.test_round2_permissions import auth, set_perms
from tests.test_round3_permissions import hr_environment, snapshot
CASES=[
 ('PUT','/config/work-schedule','config_center',default_work_schedule()),
 *[(m,p,'config_center',b) for p in ('/expense-approvals/threshold','/suppliers/purchase-approval-threshold','/suppliers/concentration-threshold') for m,b in [('GET',None),('PUT',{'threshold':80})]],
 ('GET','/rates','settings',None),('POST','/rates','settings',{'from_currency':'CNY','to_currency':'THB','rate':5}),
 ('GET','/settings/boss-contact','settings',None),('PUT','/settings/boss-contact','settings',{'name':'老板','phone':'123'}),
 ('POST','/settings/backup','settings',None),('GET','/users','settings',None),
 ('POST','/users','settings',{'username':'new-labor','password':'test-password','display_name':'新劳工','role':'warehouse_labor','warehouse_id':1,'warehouse_ids':[1]}),
 ('PUT','/users/5','settings',{'display_name':'改名'}),('DELETE','/users/5','settings',None),
 *[('GET','/efficiency/'+p,'labor_efficiency',None) for p in ('summary','trend','compare','daily-orders','standard','standards')],
 ('PUT','/efficiency/order-count','labor_efficiency',{'date':'2026-10-10','order_count':100}),
 ('PUT','/efficiency/standard','labor_efficiency',{'standard':500,'month':'2026-10'}),
 ('PUT','/efficiency/manual-hour','labor_efficiency',{'employee_id':1,'date':'2026-10-10','hours':8}),
 ('DELETE','/efficiency/manual-hour?employee_id=1&date=2026-10-10','labor_efficiency',None),
 *[('GET','/reports/'+p,'reports',None) for p in ('previews','recharge-summary','incoming-summary','income-expense','payable','expense-fund','reimbursement','credit','reconciliation-diff')],
 ('GET','/income-expense/ledger','ledger',None),('GET','/income-expense/ledger/export','ledger',None),
 ('GET','/reconciliation/unmatched','reconciliation',None),
 ('POST','/reconciliation/manual-match?declaration_id=1&flow_id=1','reconciliation',None),
 ('POST','/reconciliation/unmatch?record_id=1','reconciliation',None),
 ('GET','/reconciliation/results','reconciliation',None),('GET','/reconciliation/export','reconciliation',None),
]
@asynccontextmanager
async def r6_environment(tmp_path,monkeypatch):
 async with hr_environment(tmp_path,monkeypatch) as (engine,factory,client):
  from sqlalchemy import event
  @event.listens_for(engine.sync_engine,"connect")
  def sqlite_functions(conn,_):conn.create_function("to_char",2,lambda value,fmt:str(value)[:7] if value else None)
  await engine.dispose()
  async with factory() as db:
   db.add_all([IncomingFlow(id=i,warehouse_id=i,received_date=datetime(2026,10,10),amount=100,entrant_id=1,currency='THB',payer_name=f'仓库{i}',match_status='unmatched') for i in (1,2)])
   db.add_all([ReconciliationResult(id=i,warehouse_id=i,reconciliation_month='2026-10',amount_diff=i,match_status='manual_matched',handling_note=f'仓库{i}') for i in (1,2)])
   await db.commit()
  yield engine,factory,client
async def call(client,method,path,body,uid):
 return await client.request(method,'/api/v1'+path,headers=auth(uid),**({'json':body} if body is not None else {}))
def expected(uid,method,path):
 path=path.split('?')[0]
 if path.startswith('/settings/boss-contact'):return 200 if uid in (1,2) else 403
 if path=='/settings/backup':return 200 if uid==2 else 403
 if path=='/suppliers/concentration-threshold' and method=='GET':return 403 if uid==2 else 200
 if path in ('/config/work-schedule','/suppliers/purchase-approval-threshold','/suppliers/concentration-threshold'):
  if uid==2:return 403
  if method=='PUT':return 200 if uid==1 else 403
  return 200
 if path=='/expense-approvals/threshold':return 403 if uid==2 or (method=='PUT' and uid!=1) else 200
 if path.startswith('/rates'):return 403 if uid==2 or (method=='POST' and uid!=1) else 200
 if path.startswith('/efficiency'):
  if uid==2:return 403
  if path=='/efficiency/standard' and method=='PUT':return 200 if uid==1 else 403
  return 200
 if path=='/users' and method=='GET':return 200
 if path=='/users' and method=='POST':return 200 if uid in (1,2) else 403
 if path=='/users/5':
  if method=='DELETE':return 200 if uid==1 else (400 if uid==2 else 403)
  return 200 if uid in (1,2) else 403
 if path.startswith('/reconciliation') or path.startswith('/reports') or path.startswith('/income-expense/ledger'):return 200 if uid in (1,3) else 403
 return 403 if uid==2 else 200
@pytest.mark.parametrize('method,path,module,body',CASES,ids=[f'{m} {p}' for m,p,_,_ in CASES])
@pytest.mark.parametrize('uid,grant',[(3,False),(4,False),(3,True),(4,True),(1,False),(2,False)],ids=['staff-deny','supervisor-deny','staff-grant','supervisor-grant','admin','super-admin'])
def test_endpoint_matrix(tmp_path,monkeypatch,method,path,module,body,uid,grant):
 async def run():
  async with r6_environment(tmp_path,monkeypatch) as (_,factory,client):
   if grant:await set_perms(factory,uid,[MENU_PERMISSIONS_BY_KEY[module].label])
   before=await snapshot(factory)
   response=await call(client,method,path,body,uid)
   if uid in (3,4) and not grant:
    assert response.status_code==403,response.text
    assert 'detail' in response.json() and 'data' not in response.json(),response.text
    assert await snapshot(factory)==before
   else:
    assert response.status_code==expected(uid,method,path),response.text
    if response.status_code==403:assert 'detail' in response.json() and 'data' not in response.json(),response.text
    if path=='/settings/backup' and uid!=2:assert response.status_code==403 and 'data' not in response.json()
    if path!='/users/5':assert await snapshot(factory)==before
 asyncio.run(run())
@pytest.mark.parametrize('report',[p for m,p,k,b in CASES if k=='reports' and p!='/reports/previews'])
def test_report_excel(tmp_path,monkeypatch,report):
 async def run():
  async with r6_environment(tmp_path,monkeypatch) as (_,factory,client):
   await set_perms(factory,3,['报表中心'])
   response=await client.get('/api/v1'+report+'?format=excel',headers=auth(3))
   assert response.status_code==200,response.text
   from openpyxl import load_workbook
   from io import BytesIO
   assert load_workbook(BytesIO(response.content)).active.max_row>=1
 asyncio.run(run())
def test_route_inventory():
 root=Path(__file__).resolve().parents[1]/'app/api/v1';seen=set()
 for file,prefix in [('efficiency','efficiency'),('reports','reports'),('reconciliation','reconciliation')]:
  for n in ast.parse((root/(file+'.py')).read_text()).body:
   if isinstance(n,ast.AsyncFunctionDef):
    for d in n.decorator_list:
     if isinstance(d,ast.Call) and isinstance(d.func,ast.Attribute):seen.add((d.func.attr.upper(),'/'+prefix+d.args[0].value))
 assert seen=={(m,p.split('?')[0]) for m,p,k,b in CASES if k in ('labor_efficiency','reports','reconciliation')} | {('GET','/reports/operations')}
def test_migration_round6(tmp_path,monkeypatch):
 async def run():
  async with r6_environment(tmp_path,monkeypatch) as (engine,factory,_):
   await set_perms(factory,3,None);await set_perms(factory,4,['操作日志','配置中心','配置中心'])
   async with factory() as db:
    db.add(User(id=7,username='inactive',display_name='停用财务',role='staff',password_hash='unused',is_active=False,extra_permissions=[]))
    db.add_all([SystemSetting(warehouse_id=1,key=getattr(migrations,f'ROUND{r}_PERMISSION_MARKER'),value='true') for r in (2,3,4,5)])
    db.add(SystemSetting(warehouse_id=1,key=migrations.ROUND9_SUPERVISOR_PERMISSION_MARKER,value='true'))
    db.add(SystemSetting(warehouse_id=1,key=migrations.MODIFICATION_LOGS_PERMISSION_MARKER,value='true'))
    db.add(SystemSetting(warehouse_id=1,key=migrations.CASH_ADVANCES_PERMISSION_MARKER,value='true'))
    await db.commit()
   before=await snapshot(factory)
   with pytest.raises(RuntimeError):
    async with engine.begin() as conn:
     await migrations.backfill_round6_permissions(conn);raise RuntimeError('rollback')
   async with factory() as db:
    assert (await db.get(User,3)).extra_permissions is None
    assert not (await db.execute(select(SystemSetting).where(SystemSetting.key==migrations.ROUND6_PERMISSION_MARKER))).first()
   async with engine.begin() as conn:assert await migrations.run_permission_migrations(conn)=={'round2':0,'round3':0,'round4':0,'round5':0,'round6':1,'round9':0,'modification_logs':0,'cash_advances':0}
   async with factory() as db:
    assert (await db.get(User,3)).extra_permissions is None
    assert (await db.get(User,4)).extra_permissions==['操作日志',*migrations.ROUND6_PERMISSIONS]
    assert (await db.get(User,7)).extra_permissions==[]
    assert (await db.get(User,5)).extra_permissions==[]
    assert len((await db.execute(select(SystemSetting).where(SystemSetting.key==migrations.ROUND6_PERMISSION_MARKER))).scalars().all())==1
   await set_perms(factory,4,[])
   async with engine.begin() as conn:assert await migrations.backfill_round6_permissions(conn)==0
   async with factory() as db:assert (await db.get(User,4)).extra_permissions==[]
   assert await snapshot(factory)==before
 asyncio.run(run())

@pytest.mark.parametrize('uid',[3,4])
def test_shared_reads_and_cockpit_unchanged(tmp_path,monkeypatch,uid):
 async def run():
  async with r6_environment(tmp_path,monkeypatch) as (engine,factory,client):
   from sqlalchemy import event
   # SQLite test adapter for the PostgreSQL month formatter; business queries stay real.
   @event.listens_for(engine.sync_engine,'connect')
   def add_functions(conn,_):conn.create_function('to_char',2,lambda value,fmt:str(value)[:7] if value else None)
   await engine.dispose()
   paths=['/dashboard/stats','/dashboard/pending-tasks','/dashboard/trends','/dashboard/cockpit']
   before=await snapshot(factory)
   for p in paths:
    r=await client.get('/api/v1'+p,headers=auth(uid))
    assert r.status_code==(200 if p == '/dashboard/stats' else 403),r.text
   await set_perms(factory,uid,list(migrations.ROUND6_PERMISSIONS))
   for p in paths:
    r=await client.get('/api/v1'+p,headers=auth(uid))
    assert r.status_code==(200 if uid==3 or p == '/dashboard/stats' else 403),r.text
   assert await snapshot(factory)==before
   await set_perms(factory,uid,[])
   r=await client.get('/api/v1/users/permissions',headers=auth(uid));assert r.status_code==200
   from app.core.menu_registry import MENU_PERMISSIONS
   assert len(r.json())==len(MENU_PERMISSIONS)
   for key in ('config_center','settings','labor_efficiency','reports','ledger','reconciliation'):
    entry=next(x for x in r.json() if x['menu_key']==key);assert entry['selectable'];assert 'supervisor' not in entry['default_roles'];assert 'staff' not in entry['default_roles']
   r=await client.get('/api/v1/config/work-schedule',headers=auth(uid));assert r.status_code==403
   # Other consuming pages work with their own permission and without config/settings.
   for label,path in [('费用审批','/expense-approvals/threshold'),('供应商管理','/suppliers/purchase-approval-threshold'),('采购分析','/suppliers/concentration-threshold')]:
    await set_perms(factory,uid,[label]);r=await client.get('/api/v1'+path,headers=auth(uid))
    assert r.status_code==(403 if uid==3 and label=='供应商管理' else 200),r.text
   r=await client.get('/api/v1/rates/query',headers=auth(uid));assert r.status_code==403 and set(r.json())=={'detail'}
 asyncio.run(run())

@pytest.mark.parametrize('uid',[1,3,4])
def test_warehouse_boundaries(tmp_path,monkeypatch,uid):
 async def run():
  async with r6_environment(tmp_path,monkeypatch) as (_,factory,client):
   await set_perms(factory,uid,list(migrations.ROUND6_PERMISSIONS))
   for path in ('/reconciliation/unmatched?warehouse_id=2','/reconciliation/results?warehouse_id=2','/reconciliation/export?warehouse_id=2'):
    r=await client.get('/api/v1'+path,headers=auth(uid));assert r.status_code==403,r.text
   r=await client.post('/api/v1/reconciliation/unmatch?record_id=2',headers=auth(uid));assert r.status_code==403,r.text
   r=await client.post('/api/v1/reconciliation/manual-match?declaration_id=1&flow_id=2',headers=auth(uid));assert r.status_code==403,r.text
   for path in ('/reports/incoming-summary?warehouse_id=2','/reports/reconciliation-diff?warehouse_id=2'):
    r=await client.get('/api/v1'+path,headers=auth(uid));assert r.status_code==403 and set(r.json())=={'detail'},r.text
   r=await client.put('/api/v1/efficiency/manual-hour',headers=auth(uid),json={'employee_id':2,'date':'2026-10-10','hours':8});assert r.status_code==400,r.text
   # Forged warehouse selection must not move reads or writes to foreign warehouse.
   h={**auth(uid),'X-Warehouse-ID':'2'}
   if uid!=3:
    r=await client.put('/api/v1/efficiency/order-count',headers=h,json={'date':'2026-10-10','order_count':123});assert r.status_code==403,r.text
    from app.models.labor_efficiency import EfficiencyOrderCount
    async with factory() as db:
     rows=(await db.execute(select(EfficiencyOrderCount))).scalars().all();assert len(rows)==0
   r=await client.get('/api/v1/reconciliation/results',headers=h);assert r.status_code==403 and set(r.json())=={'detail'},r.text
 asyncio.run(run())

@pytest.mark.parametrize("setter",[1,4])
def test_config_redlines_clock_and_payroll_without_config_permission(tmp_path,monkeypatch,setter):
 async def run():
  from tests.test_round4_permissions import attendance_environment, call as attendance_call, CASES as ATTENDANCE_CASES
  from tests.test_round2_permissions import png
  from app.services.schedule import get_work_schedule
  async with attendance_environment(tmp_path,monkeypatch) as (_,factory,client):
   settings={**default_work_schedule(),'morning_start':'08:30','late_half':'08:35','late_one':'08:50','overtime_limit':65}
   before=await snapshot(factory)
   await set_perms(factory,setter,['配置中心'])
   r=await client.put('/api/v1/config/work-schedule',headers=auth(setter),json=settings)
   if setter==4:
    assert r.status_code==403 and set(r.json())=={'detail'}
    async with factory() as db:assert await get_work_schedule(db,1)==default_work_schedule()
    r=await client.put('/api/v1/config/work-schedule',headers=auth(1),json=settings)
   assert r.status_code==200,r.text
   assert await snapshot(factory)==before
   async with factory() as db:
    schedule=await get_work_schedule(db,1);assert all(schedule[k]==v for k,v in settings.items())
    foreign=await get_work_schedule(db,2);assert foreign==default_work_schedule()
   r=await attendance_call(client,*ATTENDANCE_CASES[0][:2],ATTENDANCE_CASES[0][3],5)
   assert r.status_code==200,r.text
   assert r.json()['status']=='late_one',r.text
   r=await client.get('/api/v1/clock-in/today',headers=auth(5));assert r.status_code==200,r.text
   assert '08:30' in str(r.json()),r.text
   await set_perms(factory,3,['工资'])
   r=await client.post('/api/v1/payroll/preview',headers=auth(3),json={'end_date':'2026-10-15','employee_ids':[1]});assert r.status_code==200,r.text
   assert await snapshot(factory)==before
   r=await client.post('/api/v1/payroll/calculate',headers=auth(3),json={'end_date':'2026-10-15','employee_ids':[4]});assert r.status_code==200,r.text
   after=await snapshot(factory);assert all(row in after for row in before)
   assert r.json()['record_count']==1,r.text
 asyncio.run(run())

@pytest.mark.parametrize('uid',[1,3,4])
def test_config_settings_ledger_warehouse_scope(tmp_path,monkeypatch,uid):
 async def run():
  from app.models.recharge import ExchangeRate
  async with r6_environment(tmp_path,monkeypatch) as (_,factory,client):
   await set_perms(factory,uid,list(migrations.ROUND6_PERMISSIONS))
   h={**auth(uid),'X-Warehouse-ID':'2'}
   r=await client.get('/api/v1/income-expense/ledger',headers=h)
   assert r.status_code==403 and set(r.json())=={'detail'},r.text
   r=await client.get('/api/v1/users',headers=h)
   assert r.status_code==403 and set(r.json())=={'detail'},r.text
   r=await client.post('/api/v1/users',headers=h,json={'username':'cross-wh','password':'test-password','display_name':'越仓','role':'warehouse_labor','warehouse_ids':[2]})
   assert r.status_code==403,r.text
   if uid==1:
    for method in ('put','delete'):
     r=await getattr(client,method)('/api/v1/users/6',headers=h,**({'json':{'display_name':'越仓'}} if method=='put' else {}));assert r.status_code==403,r.text
    r=await client.put('/api/v1/config/work-schedule',headers=h,json={**default_work_schedule(),'overtime_limit':75});assert r.status_code==403,r.text
    r=await client.post('/api/v1/rates',headers=h,json={'from_currency':'CNY','to_currency':'THB','rate':5});assert r.status_code==403,r.text
    async with factory() as db:
     rows=(await db.execute(select(ExchangeRate))).scalars().all();assert len(rows)==0
     settings=(await db.execute(select(SystemSetting))).scalars().all();assert not settings
 asyncio.run(run())
