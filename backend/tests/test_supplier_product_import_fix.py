"""Physical Excel files -> multipart HTTP -> disposable native PostgreSQL only."""
import asyncio
import io
import json
import os
from pathlib import Path
import uuid

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, func, text
from app.main import app
from app.models import SupplierProduct, ProductImportOperation, Supplier, PurchaseOrder
from tests.test_supplier_sku_round5 import fixture
from tests.test_supplier_workflow_round1 import request
from tests.test_round2_permissions import auth
from tests.test_incoming_operations import URL

pytestmark = pytest.mark.skipif(not URL, reason='Explicit isolated localhost PostgreSQL required')
ACTUAL = {}
HEADERS = ['供应商名称', '产品名', '产品规格', '规格报价', '单价', '单位', '币种']


def evidence(name, value):
    ACTUAL[name] = value
    if os.environ.get('SUPPLIER_IMPORT_REPORT_DIR'):
        folder = Path(os.environ['SUPPLIER_IMPORT_REPORT_DIR'])
        folder.mkdir(parents=True, exist_ok=True)
        (folder / 'actual.json').write_text(json.dumps(ACTUAL, ensure_ascii=False, default=str, indent=2))


def excel_file(tmp_path, name, rows, headers=HEADERS):
    wb = Workbook()
    if headers:
        wb.active.append(headers)
    for row in rows:
        wb.active.append(row)
    path = tmp_path / name
    wb.save(path)
    if os.environ.get('SUPPLIER_IMPORT_REPORT_DIR'):
        folder = Path(os.environ['SUPPLIER_IMPORT_REPORT_DIR']) / 'files'
        folder.mkdir(parents=True, exist_ok=True)
        (folder / name).write_bytes(path.read_bytes())
    return path


async def upload(client, path, file, **kwargs):
    with file.open('rb') as stream:
        return await request(client, 'POST', path, files={'file': (file.name, stream)}, **kwargs)


def confirm(preview, key=None):
    return {'mode': 'confirm', 'request_key': key or uuid.uuid4().hex,
            'preview_version': preview['preview_version'],
            'choices': json.dumps([{'row_number': r['row_number'], 'decision': 'unmapped', 'confirmed': True}
                                  for r in preview['rows'] if not r.get('skipped') and not r['errors']])}


async def count(factory):
    async with factory() as db:
        return (await db.execute(select(func.count(SupplierProduct.id)))).scalar_one()


def test_fixed_prefixes_are_registered_first_and_global_upload_is_reachable(tmp_path, monkeypatch):
    async def run():
        async with fixture(tmp_path, monkeypatch) as (_, factory, client, _):
            file = excel_file(tmp_path, 'route-single.xlsx', [['本仓供应商', '路由袋', 'A', None, 9, '个', 'THB']])
            result = await upload(client, '/suppliers/import/products', file)
            evidence('global_route', result)
            assert result['valid'] and await count(factory) == 1
            routes = [r for r in app.routes if getattr(r, 'path', '').startswith('/api/v1/suppliers')]
            literals = [i for i, r in enumerate(routes) if not r.path.startswith('/api/v1/suppliers/{supplier_id}')]
            suppliers = [i for i, r in enumerate(routes) if r.path.startswith('/api/v1/suppliers/{supplier_id}')]
            assert max(literals) < min(suppliers)
    asyncio.run(run())


def test_physical_xlsx_single_supplier_append_quote_and_exact_retry(tmp_path, monkeypatch):
    async def run():
        async with fixture(tmp_path, monkeypatch) as (_, factory, client, _):
            async with factory() as db:
                old = dict((await db.execute(select(SupplierProduct.__table__).where(SupplierProduct.id == 1))).mappings().one())
                old_order = (await db.get(PurchaseOrder, 1)).items
            file = excel_file(tmp_path, 'single-existing-quote.xlsx', [['本仓供应商', '耗材', None, None, 90, '个', 'THB']])
            pre = await upload(client, '/suppliers/1/import/products', file)
            assert pre['valid'] and pre['total_rows'] == pre['valid_count'] == 1 and pre['skipped_count'] == 0
            payload = confirm(pre, 'append-existing-quote')
            result = await upload(client, '/suppliers/1/import/products', file, data=payload)
            retry = await upload(client, '/suppliers/1/import/products', file, data=payload)
            assert result['imported'] == 1 and result['skipped'] == result['updated'] == 0
            assert '导入成功 1 条' in result['message'] and retry['replayed']
            assert await count(factory) == 2
            async with factory() as db:
                assert dict((await db.execute(select(SupplierProduct.__table__).where(SupplierProduct.id == 1))).mappings().one()) == old
                assert (await db.get(PurchaseOrder, 1)).items == old_order
                assert (await db.get(SupplierProduct, result['data'][0]['product_id'])).unit_price == 90
            evidence('single_append_and_retry', {'preview': pre, 'confirmed': result, 'retry': retry, 'old_quote_and_order_unchanged': True})
    asyncio.run(run())


def test_multisupplier_unknown_skip_and_single_supplier_whole_batch_rejection(tmp_path, monkeypatch):
    async def run():
        async with fixture(tmp_path, monkeypatch) as (_, factory, client, _):
            file = excel_file(tmp_path, 'multiple-suppliers.xlsx', [
                ['本仓供应商', '甲袋', 'A', None, 9, '个', 'THB'],
                ['同开票方别名', '乙袋', 'B', None, 8, '个', 'THB'],
                ['未建档供应商', '不猜配', 'C', None, 7, '个', 'THB'],
                ['## 示例行，不会导入', '示例袋', 'D', None, 1, '个', 'THB']])
            pre = await upload(client, '/suppliers/import/products', file)
            assert pre['valid'] and (pre['total_rows'], pre['valid_count'], pre['skipped_count']) == (4, 2, 2)
            assert [r['row_number'] for r in pre['skipped_rows']] == [4, 5]
            single = await upload(client, '/suppliers/1/import/products', file)
            assert not single['valid'] and any(e['row_number'] == 3 and '供应商' in e['message'] for e in single['errors'])
            rejected = await upload(client, '/suppliers/1/import/products', file, code=422, data=confirm(single))
            assert await count(factory) == 1
            result = await upload(client, '/suppliers/import/products', file, data=confirm(pre))
            assert (result['imported'], result['skipped'], result['updated']) == (2, 2, 0)
            assert await count(factory) == 3
            async with factory() as db:
                created = (await db.execute(select(SupplierProduct).where(SupplierProduct.id != 1))).scalars().all()
                assert {p.supplier_id for p in created} == {1, 3}
                assert all(p.supplier_id != 2 for p in created)
            evidence('multiple_and_single', {'preview': pre, 'single_preview': single, 'single_rejected': rejected, 'global_confirmed': result})
    asyncio.run(run())


def test_downloaded_original_template_is_example_only_and_empty_headers_are_chinese(tmp_path, monkeypatch):
    async def run():
        async with fixture(tmp_path, monkeypatch) as (_, factory, client, _):
            r = await client.get('/api/v1/suppliers/import-template/products', headers=auth(1))
            assert r.status_code == 200
            path = tmp_path / 'original-template.xlsx'
            path.write_bytes(r.content)
            wb = load_workbook(path)
            assert str(wb.active['A2'].value).startswith('##')
            pre = await upload(client, '/suppliers/import/products', path)
            assert not pre['valid'] and pre['total_rows'] == pre['skipped_count'] == 1 and pre['valid_count'] == 0
            assert '全部是示例行' in pre['message'] and '没有可导入数据' in pre['message']
            empty = tmp_path / 'empty.xlsx'; empty.write_bytes(b'')
            no_bytes = await upload(client, '/suppliers/import/products', empty, code=400)
            no_header = excel_file(tmp_path, 'no-header.xlsx', [], headers=None)
            no_header_result = await upload(client, '/suppliers/import/products', no_header, code=400)
            assert '文件为空' in no_bytes['detail'] and '表头' in no_header_result['detail']
            assert await count(factory) == 1
            evidence('original_template_empty_and_headers', {'preview': pre, 'empty': no_bytes, 'no_header': no_header_result})
    asyncio.run(run())


def test_xls_has_explicit_conversion_message_and_bad_rows_stay_atomic(tmp_path, monkeypatch):
    async def run():
        async with fixture(tmp_path, monkeypatch) as (_, factory, client, _):
            xls = Path(__file__).parent / 'fixtures' / 'products_legacy.xls'
            response = await upload(client, '/suppliers/import/products', xls, code=400)
            assert response['detail'] == '请另存为 xlsx 格式'
            file = excel_file(tmp_path, 'invalid-price.xlsx', [
                ['本仓供应商', '正常', 'A', None, 9, '个', 'THB'],
                ['本仓供应商', '错误', 'B', None, '=1+1', '个', 'THB']])
            pre = await upload(client, '/suppliers/import/products', file)
            assert not pre['valid']
            await upload(client, '/suppliers/import/products', file, code=422, data=confirm(pre))
            assert await count(factory) == 1
            evidence('xls_and_invalid_atomic', {'xls': response, 'invalid': pre})
    asyncio.run(run())


def test_skipped_supplier_cannot_be_forged_or_silently_included_after_preview(tmp_path, monkeypatch):
    async def run():
        async with fixture(tmp_path, monkeypatch) as (_, factory, client, _):
            file = excel_file(tmp_path, 'unknown-empty-product.xlsx', [
                ['本仓供应商', '核对袋', 'A', None, 9, '个', 'THB'],
                ['尚未建档', None, None, None, '=1+1', '个', 'THB'],
                ['等待建档', '未来可用报价', 'C', None, 8, '个', 'THB']])
            pre = await upload(client, '/suppliers/import/products', file)
            assert pre['valid'] and pre['valid_count'] == 1 and pre['skipped_count'] == 2
            data = confirm(pre)
            data['choices'] = json.dumps([{'row_number': number, 'decision': 'unmapped', 'confirmed': True} for number in (2, 3, 4)])
            rejected = await upload(client, '/suppliers/import/products', file, code=400, data=data)
            assert await count(factory) == 1
            async with factory() as db:
                # The supplied native fixture uses explicit supplier ids too.
                await db.execute(text("SELECT setval(pg_get_serial_sequence('suppliers','id'),(SELECT max(id) FROM suppliers))"))
                db.add(Supplier(warehouse_id=1, name='等待建档'))
                await db.commit()
            changed = await upload(client, '/suppliers/import/products', file, code=409, data=confirm(pre))
            assert changed['detail']['code'] == 'product_import_preview_changed'
            assert await count(factory) == 1
            evidence('skipped_choice_and_changed_supplier', {'preview': pre, 'forged_choice': rejected, 'changed_preview': changed})
    asyncio.run(run())


def test_existing_permissions_and_warehouse_rejections_remain(tmp_path, monkeypatch):
    async def run():
        async with fixture(tmp_path, monkeypatch) as (_, factory, client, _):
            file = excel_file(tmp_path, 'permissions.xlsx', [['本仓供应商', '权限袋', 'A', None, 9, '个', 'THB']])
            for uid in (2, 5, 8, 10):
                await upload(client, '/suppliers/import/products', file, uid=uid, code=403)
            foreign = await upload(client, '/suppliers/2/import/products', file, code=404)
            assert await count(factory) == 1
            allowed = await upload(client, '/suppliers/import/products', file, uid=9)
            assert allowed['valid'] and not allowed['rows'][0]['catalog_visible']
            assert not allowed['rows'][0]['candidates']
            evidence('permissions', {'super_labor_ungranted_catalog_only': 403, 'foreign': foreign,
                                     'supplier_only_preview': allowed, 'products_unchanged': True})
    asyncio.run(run())
