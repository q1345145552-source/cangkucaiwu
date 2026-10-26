from app.services.permission_migrations import MODIFICATION_LOGS_PERMISSION_MARKER, CASH_ADVANCES_PERMISSION_MARKER
"""第五轮真实HTTP：38条接口、财务状态、补权、共用读取和仓库隔离。"""
import asyncio, ast, json, re
from contextlib import asynccontextmanager
from datetime import datetime
from pathlib import Path
import pytest
from sqlalchemy import select, update, func
from app.models import User, SystemSetting
from app.models.supplier import Supplier, SupplierCategory, PurchaseOrder
from app.models.payable import PayableBill, PayablePlan, PlanTemplate, MonthlyOrderVolume
from app.models.group_order import GroupOrder, GroupOrderParticipant
from app.core.timezone import THAI_TZ
from app.core.menu_registry import MENU_PERMISSIONS_BY_KEY
from app.services.permission_migrations import ROUND2_PERMISSION_MARKER, ROUND3_PERMISSION_MARKER, ROUND4_PERMISSION_MARKER, ROUND5_PERMISSION_MARKER, ROUND6_PERMISSION_MARKER, ROUND9_SUPERVISOR_PERMISSION_MARKER, ROUND5_PERMISSIONS, backfill_round5_permissions, run_permission_migrations
from tests.test_round2_permissions import environment, auth, set_perms, png

CASES=[
 ('GET','/payable','payable',None),
 ('POST','/payable','payable',{'form':{'supplier_id':'1','bill_number':'NEW-BILL','bill_date':'2026-10-10','due_date':'2026-10-20','amount':'100'},'file':True}),
 ('POST','/payable/1/upload-voucher','payable',{'file':True}),
 ('POST','/payable/1/upload-attachment','payable',{'file':True}),
 ('PUT','/payable/1/pay?paid_amount=100&payment_method=bank','payable',None),
 ('PUT','/payable/4/confirm','payable',None),
 ('PUT','/payable/1/void?void_reason=测试作废','payable',None),
 ('PUT','/payable/1','payable',{'remark':'已编辑','amount':120}),
 ('GET','/payable/stats?start_date=2026-10-01&end_date=2026-10-31','payable',None),
 ('GET','/payable/batch-export','payable',None),
 ('GET','/payable/supplier-statement','payable',None),
 ('GET','/payable/monthly-order','payable',None),
 ('POST','/payable/monthly-order','payable',{'month':'2026-10','order_count':100}),
 ('GET','/payable/trend?months=1','payable',None),
 ('GET','/payable/plans','payment_plans',None),
 ('POST','/payable/plans','payment_plans',{'plan_name':'新计划','planned_date':'2026-10-20','bill_ids':[1,2,3,5,6,7],'save_as_template':True}),
 ('GET','/payable/cashflow-prediction','payment_plans',None),
 ('GET','/payable/plan-templates','payment_plans',None),
 ('DELETE','/payable/plan-templates/1','payment_plans',None),
 ('GET','/payable/timeline','payment_plans',None),
 ('PUT','/payable/plans/1/execute','payment_plans',None),
 ('GET','/payable/plans/1/detail','payment_plans',None),
 ('POST','/payable/plans/1/upload-attachment','payment_plans',{'file':True}),
 ('GET','/payable/plans/export','payment_plans',None),
 ('GET','/suppliers/concentration?month=2026-10','concentration',None),
 ('GET','/suppliers/concentration-trend?months=1','concentration',None),
 ('GET','/suppliers/concentration-threshold','concentration',None),
 ('PUT','/suppliers/concentration-threshold','concentration',{'threshold':80}),
 ('GET','/group-order','group_order',None),
 ('POST','/group-order','group_order',{'item_name':'新商品','target_quantity':100,'target_price':10,'deadline':'2026-11-01'}),
 ('POST','/group-order/10/join','group_order',{'quantity':10,'delivery_address':'测试仓','agreed_rules':True}),
 ('GET','/group-order/1/participants','group_order',None),
 ('GET','/group-order/history','group_order',None),
 ('GET','/group-order/1','group_order',None),
 ('PUT','/group-order/1/close','group_order',None),
 ('PUT','/group-order/2/complete','group_order',{'final_price':9,'final_supplier':'测试商','logistics_fee':10}),
 ('PUT','/group-order/1/cancel','group_order',{'reason':'测试取消'}),
 ('PUT','/group-order/participant/1/ban','group_order',None),
]

@asynccontextmanager
async def finance_environment(tmp_path,monkeypatch):
    async with environment(tmp_path,monkeypatch) as (engine,factory,client):
        import app.api.v1.payable as payable
        import app.api.v1.suppliers as suppliers
        import app.services.image_utils as images
        fixed=datetime(2026,10,16,9,tzinfo=THAI_TZ)
        for module in (payable,suppliers):
            monkeypatch.setattr(module,'thai_today',lambda:fixed.date())
            monkeypatch.setattr(module,'thai_now',lambda:fixed)
        original_save=images.save_image
        def save_in_test(content,absolute,relative,filename):
            return original_save(content,str(tmp_path/'finance-uploads'/absolute.lstrip('/')),relative,filename)
        monkeypatch.setattr(images,'save_image',save_in_test)
        async with factory() as db:
            db.add(SupplierCategory(id=1,name='耗材商'))
            db.add_all([Supplier(id=i,warehouse_id=i,name=f'供应商{i}',category_id=1) for i in (1,2)])
            for po_id,wid,status,amount in ((1,1,'confirmed',700),(2,1,'pending',600),(10,2,'confirmed',10000)):
                db.add(PurchaseOrder(id=po_id,warehouse_id=wid,supplier_id=wid,order_number=f'PO-{po_id}',
                                     total_amount=amount,currency='THB',items=[],status=status,created_at=datetime(2026,10,10)))
            for bid,wid,amount,paid,status,po,confirm in [
                (1,1,100,0,'pending',None,'false'),(2,1,200,200,'paid',None,'false'),
                (3,1,9000,0,'voided',None,'false'),(4,1,400,0,'pending',None,'true'),
                (5,1,150,50,'partially_paid',None,'false'),(6,1,600,0,'pending',2,'false'),
                (7,1,700,0,'pending',1,'false'),(10,2,10000,0,'pending',None,'false'),
            ]:
                db.add(PayableBill(id=bid,warehouse_id=wid,supplier_id=wid,bill_number=f'BILL-{bid}',
                                   bill_date=datetime(2026,10,10),due_date=datetime(2026,10,10 if wid==2 else 20),
                                   amount=amount,paid_amount=paid,status=status,currency='THB',purchase_order_id=po,
                                   need_boss_confirm=confirm,created_by=1))
            for pid,wid,ids in ((1,1,[1,2,3,5]),(10,2,[10])):
                db.add(PayablePlan(id=pid,warehouse_id=wid,plan_name=f'计划{pid}',planned_date=datetime(2026,10,20),
                                   total_amount=200 if wid==1 else 10000,bill_ids=ids,status='pending',created_by=1))
                db.add(PlanTemplate(id=pid,warehouse_id=wid,name=f'模板{pid}',bill_ids=ids,created_by=1))
                db.add(MonthlyOrderVolume(warehouse_id=wid,month='2026-10',order_count=100 if wid==1 else 999))
            for gid,wid,status in ((1,1,'open'),(2,1,'closed'),(3,1,'completed'),(10,2,'open')):
                db.add(GroupOrder(id=gid,warehouse_id=wid,item_name=f'商品{gid}',target_quantity=100,target_price=10,
                                  deadline=datetime(2026,11,1),initiator_id=1,status=status))
            db.add_all([GroupOrderParticipant(id=i,group_order_id=i,warehouse_id=wid,quantity=10,agreed_rules=True)
                        for i,wid in ((1,1),(10,2))])
            await db.commit()
        yield engine,factory,client

async def finance_snapshot(factory):
    async with factory() as db:
        return {model.__tablename__:[tuple(r) for r in (await db.execute(select(model.__table__).order_by(model.id))).all()]
                for model in (PayableBill,PayablePlan,PlanTemplate,MonthlyOrderVolume,PurchaseOrder,Supplier,GroupOrder,GroupOrderParticipant)}

async def call(client,method,path,body,uid):
    kwargs={'headers':auth(uid)}
    if isinstance(body,dict) and ('form' in body or body.get('file')):
        kwargs['data']=body.get('form',{})
        if body.get('file'):kwargs['files']={'file':('test.png',png(),'image/png')}
    elif body is not None:kwargs['json']=body
    return await client.request(method,'/api/v1'+path,**kwargs)

def expected_status(uid,method,path):
    path=path.split('?')[0]
    if method=='GET' and path.startswith('/suppliers/concentration'):return 200
    if path=='/group-order/participant/1/ban':return 200 if uid==2 else 403
    if uid==2:
        return 200 if path in ('/payable/plan-templates/1','/payable/plans/1/detail','/group-order/1/participants',
                               '/group-order/1','/group-order/1/close','/group-order/2/complete','/group-order/1/cancel') else 403
    if path=='/payable/4/confirm':return 200 if uid==1 else 403
    if path=='/suppliers/concentration-threshold' and method=='PUT':return 200 if uid==1 else 403
    if uid==3:
        if path.startswith('/suppliers/concentration'):return 403
        if path.startswith('/payable'):
            if (method=='POST' and path=='/payable/plans') or method=='DELETE' or path.endswith('/execute'):return 200
            return 200 if method=='GET' else 403
        if path.startswith('/group-order'):
            return 200 if method=='GET' or path.endswith('/join') else 403
    return 200

@pytest.mark.parametrize('method,path,module,body',CASES,ids=[f'{m} {p}' for m,p,_,_ in CASES])
@pytest.mark.parametrize('uid,granted',[(3,False),(4,False),(3,True),(4,True),(1,False),(2,False)],
                         ids=['staff-deny','supervisor-deny','staff-allow','supervisor-allow','admin','super-admin'])
def test_all_finance_endpoints(tmp_path,monkeypatch,method,path,module,body,uid,granted):
    async def run():
        async with finance_environment(tmp_path,monkeypatch) as (_,factory,client):
            if granted:await set_perms(factory,uid,[MENU_PERMISSIONS_BY_KEY[module].label])
            before=await finance_snapshot(factory)
            response=await call(client,method,path,body,uid)
            if uid in (3,4) and not granted:
                assert response.status_code==403,response.text
                assert response.json()['detail'].startswith('无此操作权限'),response.text
                assert await finance_snapshot(factory)==before
            else:
                expected=expected_status(uid,method,path)
                assert response.status_code==expected,response.text
                if expected==403:
                    if uid in (3,4) and method=='PUT' and path=='/suppliers/concentration-threshold':
                        assert response.json()['detail'].startswith('无此操作权限') and '配置中心' in response.json()['detail'],response.text
                    else:assert not response.json()['detail'].startswith('无此操作权限'),response.text
                if expected!=200 or method=='GET':assert await finance_snapshot(factory)==before
                if expected==200 and isinstance(body,dict) and body.get('file') and path!='/payable':assert response.json()['path']
    asyncio.run(run())

def test_finance_route_inventory():
    root=Path(__file__).resolve().parents[1]/'app/api/v1'
    seen=[]
    for filename,prefix in {'payable':'payable','group_order':'group-order','suppliers':'suppliers'}.items():
        for node in ast.walk(ast.parse((root/(filename+'.py')).read_text())):
            if not isinstance(node,ast.AsyncFunctionDef):continue
            for d in node.decorator_list:
                if isinstance(d,ast.Call) and isinstance(d.func,ast.Attribute) and getattr(d.func.value,'id',None)=='router':
                    path=d.args[0].value
                    if filename=='suppliers' and not path.startswith('/concentration'):continue
                    seen.append((d.func.attr.upper(),'/'+prefix+path))
                    defaults=[ast.unparse(v) for v in node.args.defaults]
                    assert any('check_staff_permission' in v for v in defaults),node.name
    assert len(seen)==38
    actual={(m,re.sub(r'/(?:1|2|4|10)(?=/|$)','/{}',p.split('?')[0])) for m,p,_,_ in CASES}
    expected={(m,re.sub(r'\{[^}]+\}','{}',p)) for m,p in seen}
    assert actual==expected

@pytest.mark.parametrize('uid',[3,4])
def test_shared_lists_and_config_threshold(tmp_path,monkeypatch,uid):
    async def run():
        async with finance_environment(tmp_path,monkeypatch) as (_,factory,client):
            for permission in ('应付账款','付款计划'):
                await set_perms(factory,uid,[permission])
                response=await client.get('/api/v1/payable',headers=auth(uid))
                assert response.status_code==200,response.text
                assert {row['warehouse_id'] for row in response.json()['data']}=={1}
                response=await client.get('/api/v1/suppliers?page_size=100',headers=auth(uid))
                assert response.status_code==200,response.text
                assert {row['id'] for row in response.json()['data']}=={1}
                # 借共用列表读取不能修改供应商，供应商原校验保持不变。
                response=await client.put('/api/v1/suppliers/1',headers=auth(uid),json={'name':'越权'})
                assert response.status_code==403 and '供应商管理' in response.json()['detail']
                wrong_path='/payable/plans/1/execute' if permission=='应付账款' else '/payable/1/pay'
                response=await client.put('/api/v1'+wrong_path,headers=auth(uid))
                assert response.status_code==403 and response.json()['detail'].startswith('无此操作权限')
            for permission in ('采购分析','配置中心'):
                await set_perms(factory,uid,[permission])
                response=await client.get('/api/v1/suppliers/concentration-threshold',headers=auth(uid))
                assert response.status_code==200,response.text
            await set_perms(factory,uid,['配置中心'])
            response=await client.get('/api/v1/suppliers/concentration?month=2026-10',headers=auth(uid))
            assert response.status_code==403 and response.json()['detail'].startswith('无此操作权限')
    asyncio.run(run())

@pytest.mark.parametrize('uid',[1,4])
def test_voided_statistics_paid_bill_edit_and_plan_execution(tmp_path,monkeypatch,uid):
    async def run():
        async with finance_environment(tmp_path,monkeypatch) as (_,factory,client):
            await set_perms(factory,uid,list(ROUND5_PERMISSIONS))
            response=await client.get('/api/v1/payable/stats?start_date=2026-10-01&end_date=2026-10-31',headers=auth(uid))
            assert response.status_code==200,response.text
            assert response.json()['month_total']==1550,response.text
            assert response.json()['month_paid']==250 and response.json()['month_unpaid']==1300
            response=await client.get('/api/v1/payable/cashflow-prediction',headers=auth(uid))
            assert response.json()['predicted_outflow']==1300,response.text
            response=await client.get('/api/v1/payable/trend?months=1',headers=auth(uid))
            assert response.json()['data'][0]['consumable_spending']==1550,response.text
            response=await client.get('/api/v1/payable/timeline',headers=auth(uid))
            assert sum(week['total'] for week in response.json()['data'])==1900,response.text
            from openpyxl import load_workbook
            from io import BytesIO
            response=await client.get('/api/v1/payable/supplier-statement',headers=auth(uid))
            rows=list(load_workbook(BytesIO(response.content)).active.values)
            assert 'BILL-3' not in {row[1] for row in rows[1:]}
            assert 'BILL-10' not in {row[1] for row in rows[1:]}
            # 已付款仍仅允许原有备注/差异处理说明白名单。
            response=await client.put('/api/v1/payable/2',headers=auth(uid),json={'remark':'新备注','diff_note':'处理说明'})
            assert response.status_code==200,response.text
            before=await finance_snapshot(factory)
            for fields in ({'amount':201},{'supplier_id':1},{'currency':'CNY'},{'bill_number':'改号'},
                           {'bill_date':'2026-10-11'},{'due_date':'2026-10-25'},{'confirmed_amount':200},
                           {'payment_commitment_days':1},{'detail':'改明细'},{'purchase_order_number':'PO-1'}):
                response=await client.put('/api/v1/payable/2',headers=auth(uid),json=fields)
                assert response.status_code==400 and response.json()['detail']=='已付款的账单只能修改备注',response.text
                assert await finance_snapshot(factory)==before
            for bid in (2,3):
                response=await client.put(f'/api/v1/payable/{bid}/void?void_reason=测试',headers=auth(uid))
                assert response.status_code==400,response.text
            # 新计划金额不包含已作废、已付清或未生效采购单的账单。
            response=await client.post('/api/v1/payable/plans',headers=auth(uid),json=CASES[15][3])
            assert response.status_code==200,response.text
            async with factory() as db:
                assert (await db.get(PayablePlan,response.json()['id'])).total_amount==900
            # 现有逐笔付款要求账户和真实可用余额；旧用例未传账户，在6a55635基线也返回400。
            from app.models import PaymentAccount
            async with factory() as db:
                account=await db.get(PaymentAccount,1)
                account.opening_balance=10000
                account.opening_balance_date=datetime(2026,10,1).date()
                await db.commit()
            response=await client.put('/api/v1/payable/plans/1/execute',headers=auth(uid),params={'account_id':1})
            assert response.status_code==200 and '2 笔' in response.json()['message'],response.text
            async with factory() as db:
                assert (await db.get(PayableBill,1)).paid_amount==100
                assert (await db.get(PayableBill,5)).paid_amount==150
                voided=await db.get(PayableBill,3)
                assert voided.status=='voided' and voided.paid_amount==0
                assert (await db.get(PayableBill,2)).paid_amount==200
            after=await finance_snapshot(factory)
            response=await client.put('/api/v1/payable/plans/1/execute',headers=auth(uid))
            assert response.status_code==400,response.text
            assert await finance_snapshot(factory)==after
    asyncio.run(run())

@pytest.mark.parametrize('uid',[1,4])
def test_void_reduces_statistics_without_deleting_bill(tmp_path,monkeypatch,uid):
    async def run():
        async with finance_environment(tmp_path,monkeypatch) as (_,factory,client):
            await set_perms(factory,uid,['应付账款'])
            response=await client.put('/api/v1/payable/1/void?void_reason=测试作废',headers=auth(uid))
            assert response.status_code==200,response.text
            async with factory() as db:
                bill=await db.get(PayableBill,1)
                assert bill.status=='voided' and bill.amount==100 and bill.void_reason=='测试作废'
            response=await client.get('/api/v1/payable/stats?start_date=2026-10-01&end_date=2026-10-31',headers=auth(uid))
            assert response.json()['month_total']==1450 and response.json()['month_paid']==250,response.text
    asyncio.run(run())

@pytest.mark.parametrize('uid',[1,4])
def test_finance_warehouse_isolation(tmp_path,monkeypatch,uid):
    async def run():
        async with finance_environment(tmp_path,monkeypatch) as (_,factory,client):
            await set_perms(factory,uid,list(ROUND5_PERMISSIONS))
            before=await finance_snapshot(factory)
            for method,path,body in [
                ('PUT','/payable/10',{'remark':'越权'}),('PUT','/payable/10/pay',None),
                ('PUT','/payable/10/confirm',None),('PUT','/payable/10/void?void_reason=越权',None),
                ('POST','/payable/10/upload-voucher',{'file':True}),('POST','/payable/10/upload-attachment',{'file':True}),
                ('POST','/payable/plans',{'plan_name':'越权计划','planned_date':'2026-10-20','bill_ids':[1,10]}),
                ('PUT','/payable/plans/10/execute',None),('GET','/payable/plans/10/detail',None),
                ('POST','/payable/plans/10/upload-attachment',{'file':True}),('DELETE','/payable/plan-templates/10',None),
                ('GET','/group-order/10/participants',None),('PUT','/group-order/10/close',None),
                ('PUT','/group-order/10/complete',{'final_price':1}),('PUT','/group-order/10/cancel',{'reason':'越权'}),
                ('PUT','/group-order/participant/10/ban',None),
            ]:
                response=await call(client,method,path,body,uid)
                assert response.status_code==403,(path,response.text)
                assert await finance_snapshot(factory)==before
            for path in ('/payable','/payable/plans','/payable/plan-templates'):
                response=await client.get('/api/v1'+path,headers={**auth(uid),'X-Warehouse-ID':'2'})
                assert response.status_code==200,response.text
                if path=='/payable':assert {row['warehouse_id'] for row in response.json()['data']}=={1}
                else:assert {row['id'] for row in response.json()['data']}=={1}
            response=await client.get('/api/v1/suppliers/concentration?month=2026-10',headers=auth(uid))
            assert response.json()['total']==700,response.text
            assert {row['supplier_id'] for row in response.json()['data']}=={1}
            # 导出外仓记录只能拿到空数据，不能越权获取。
            response=await client.get('/api/v1/payable/batch-export?bill_ids=10',headers=auth(uid))
            assert response.json()['data']==[]
            response=await client.get('/api/v1/payable/plans/export?plan_id=10',headers=auth(uid))
            from io import BytesIO
            from openpyxl import load_workbook
            assert len(list(load_workbook(BytesIO(response.content)).active.values))==1
            assert await finance_snapshot(factory)==before
            # 拼单公开浏览和跨仓参与是旧设计，仍按参与者本仓写入。
            response=await client.get('/api/v1/group-order/10',headers=auth(uid))
            assert response.status_code==200,response.text
            response=await client.post('/api/v1/group-order/10/join',headers=auth(uid),json={'quantity':5,'warehouse_id':2})
            assert response.status_code==200,response.text
            async with factory() as db:
                assert (await db.get(GroupOrderParticipant,response.json()['id'])).warehouse_id==1
    asyncio.run(run())

@pytest.mark.parametrize('path',['/payable','/payable/plans','/payable/batch-export'])
def test_overdue_updates_cannot_touch_other_warehouses(tmp_path,monkeypatch,path):
    async def run():
        async with finance_environment(tmp_path,monkeypatch) as (_,factory,client):
            await set_perms(factory,4,['应付账款','付款计划'])
            async with factory() as db:
                await db.execute(update(PayableBill).where(PayableBill.id==1).values(due_date=datetime(2026,10,10)))
                await db.commit()
            response=await client.get('/api/v1'+path,headers=auth(4))
            assert response.status_code==200,response.text
            async with factory() as db:
                assert (await db.get(PayableBill,1)).status=='overdue'
                assert (await db.get(PayableBill,10)).status=='pending'
    asyncio.run(run())

def test_round5_migration_database_idempotence_and_rollback(tmp_path,monkeypatch):
    async def run():
        async with finance_environment(tmp_path,monkeypatch) as (engine,factory,_):
            await set_perms(factory,3,None)
            await set_perms(factory,4,['供应商管理','应付账款','应付账款'])
            async with factory() as db:
                db.add(User(id=5,username='inactive-staff',display_name='停用财务',role='staff',is_active=False,password_hash='unused',extra_permissions=['旧权限']))
                db.add(User(id=6,username='labor-test',display_name='劳工',role='warehouse_labor',password_hash='unused',extra_permissions=[]))
                db.add_all([SystemSetting(warehouse_id=1,key=marker,value='true') for marker in (CASH_ADVANCES_PERMISSION_MARKER,MODIFICATION_LOGS_PERMISSION_MARKER,ROUND9_SUPERVISOR_PERMISSION_MARKER,ROUND6_PERMISSION_MARKER,ROUND2_PERMISSION_MARKER,ROUND3_PERMISSION_MARKER,ROUND4_PERMISSION_MARKER)])
                await db.commit()
            before=await finance_snapshot(factory)
            with pytest.raises(RuntimeError):
                async with engine.begin() as conn:
                    await backfill_round5_permissions(conn)
                    raise RuntimeError('模拟事务失败')
            async with factory() as db:
                assert (await db.get(User,3)).extra_permissions is None
                assert (await db.execute(select(SystemSetting.id).where(SystemSetting.key==ROUND5_PERMISSION_MARKER))).first() is None
            async with engine.begin() as conn:
                assert await run_permission_migrations(conn)=={'round2':0,'round3':0,'round4':0,'round5':3,'round6':0,'round9':0,'modification_logs':0,'cash_advances':0}
            assert await finance_snapshot(factory)==before
            async with factory() as db:
                values=dict((await db.execute(select(User.id,User.extra_permissions))).all())
                assert values[3]==list(ROUND5_PERMISSIONS)
                assert values[4]==['供应商管理',*ROUND5_PERMISSIONS]
                assert values[5]==['旧权限',*ROUND5_PERMISSIONS]
                assert values[1]==values[2]==values[6]==[]
                assert (await db.execute(select(func.count(SystemSetting.id)).where(SystemSetting.key==ROUND5_PERMISSION_MARKER))).scalar_one()==1
            await set_perms(factory,3,[])
            async with engine.begin() as conn:
                assert await run_permission_migrations(conn)=={'round2':0,'round3':0,'round4':0,'round5':0,'round6':0,'round9':0,'modification_logs':0,'cash_advances':0}
            async with factory() as db:assert (await db.get(User,3)).extra_permissions==[]
            assert await finance_snapshot(factory)==before
    asyncio.run(run())
