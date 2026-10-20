"""Native PostgreSQL: legacy money records, summary and export use the same wallet facts."""
import asyncio
import json
import os
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func, event
from app.models import (User, Customer, PaymentAccount, IncomeRecord, ExpenseRecord, IncomeExpenseCategory,
    IncomingFlow, RechargeDeclaration, ReconciliationResult, CreditCustomer, CreditRepayment, AccountTransfer,
    ExpenseFund, ExpenseFundItem, ExpenseFundDisbursement, CashAdvance, CashAdvanceRepayment,
    Reimbursement, ReimbursementPayment, Supplier, PayableBill, PayablePayment, Employee)
from app.models.payroll import PayrollRecord
from app.models.employee_advance import EmployeeAdvance
from tests.test_incoming_operations import environment, URL, body
from tests.test_round2_permissions import auth

pytestmark = pytest.mark.skipif(not URL, reason='Requires explicitly configured loopback native PostgreSQL')
RESULTS = {}
RANGE = {'start_date':'2026-10-01','end_date':'2026-10-31'}


async def read(client,path,uid=1,**params):
    r=await client.get('/api/v1/income-expense/'+path,headers=auth(uid),params={**RANGE,**params})
    assert r.status_code==200,(path,r.status_code,r.text)
    return r.json()


async def check_scope(client,kind,expected,count,uid=1,**params):
    listed=await read(client,kind,uid,**params)
    summary=await read(client,'monthly-summary',uid,**params)
    exported=await client.get('/api/v1/income-expense/'+kind+'/export',headers=auth(uid),params={**RANGE,**params})
    assert exported.status_code==200,exported.text
    sheet=list(load_workbook(BytesIO(exported.content),data_only=True).active.values)
    assert listed['total']==len(sheet)-1==count,(listed,sheet)
    assert listed['totals_by_currency']==summary[kind+'_by_currency']==expected,(listed,summary)
    by_currency={}
    for row in sheet[1:]: by_currency[row[4]]=by_currency.get(row[4],0)+row[3]
    assert by_currency==expected,(sheet,expected)
    return listed,summary,sheet


def actual(name,value):
    RESULTS[name]=value
    if os.environ.get('MONEY_RECORD_TEST_REPORT'):
        Path(os.environ['MONEY_RECORD_TEST_REPORT']).write_text(json.dumps(RESULTS,ensure_ascii=False,indent=2,default=str))


def test_real_receipt_is_visible_in_legacy_records_and_summary(tmp_path,monkeypatch):
    async def run():
        import app.services.cost_funds_debt as dates
        monkeypatch.setattr(dates,'thai_today',lambda:date(2026,10,20))
        async with environment(tmp_path,monkeypatch) as (_,factory,client):
            response=await client.post('/api/v1/incoming',headers=auth(1),json=body(remark='ACTUAL-RECEIPT-1000'))
            assert response.status_code==200,response.text
            listed=await client.get('/api/v1/income-expense/income',headers=auth(1),params=RANGE)
            summary=await client.get('/api/v1/income-expense/monthly-summary',headers=auth(1),params=RANGE)
            assert listed.status_code==summary.status_code==200
            actual('receipt_only',{'list':listed.json(),'summary':summary.json()})
            assert summary.json()['income_by_currency']=={'THB':1000}
            assert listed.json()['total']==1,listed.text
            assert listed.json()['data'][0]['amount']==1000
            assert listed.json()['data'][0]['source']=='incoming'
            assert listed.json()['data'][0]['source_link']=='/incoming'
            async with factory() as db:
                db.add(IncomeExpenseCategory(id=1,warehouse_id=1,type='income',name='业务收款',category_group='operating'))
                await db.commit()
            r=await client.post('/api/v1/income-expense/income',headers=auth(1),json={'category_id':1,'account_id':1,
                'amount':200,'income_date':'2026-10-20','currency':'THB','remark':'MANUAL-200'})
            assert r.status_code==200,r.text
            all_rows,summary,sheet=await check_scope(client,'income',{'THB':1200},2)
            pages=[]
            for p in (1,2):
                page=await read(client,'income',page=p,page_size=1)
                assert page['total']==2 and page['totals_by_currency']=={'THB':1200} and len(page['data'])==1
                pages.append(page['data'][0]['id'])
            assert len(set(pages))==2
            actual('manual_plus_receipt',{'total':summary['total_income'],'rows':all_rows['data'],'pages':pages,'export_rows':len(sheet)-1})
    asyncio.run(run())


def test_declaration_credit_links_and_internal_transfers_never_duplicate_receipts(tmp_path,monkeypatch):
    async def run():
        async with environment(tmp_path,monkeypatch) as (_,factory,client):
            async with factory() as db:
                db.add(RechargeDeclaration(id=1,warehouse_id=1,customer_id=1,declare_date=datetime(2026,10,1),amount=1000,currency='THB',declarer_id=1))
                db.add(IncomeExpenseCategory(id=1,warehouse_id=1,type='income',name='业务收款',category_group='operating'))
                db.add(PaymentAccount(id=3,warehouse_id=1,account_name='内部银行',account_type='bank',account_number='3',currency='THB'))
                await db.commit()
            await check_scope(client,'income',{},0)
            r=await client.post('/api/v1/incoming',headers=auth(1),json=body(remark='ONE-REAL-RECEIPT'))
            assert r.status_code==200;r_id=r.json()['id']
            async with factory() as db:
                db.add(CreditCustomer(id=1,warehouse_id=1,customer_id=1,credit_limit=10000))
                db.add_all([ReconciliationResult(warehouse_id=1,reconciliation_month='2026-10',declaration_id=1,flow_id=r_id,match_status='matched'),
                    ReconciliationResult(warehouse_id=1,reconciliation_month='2026-10',declaration_id=1,flow_id=r_id,match_status='manual_matched')])
                await db.flush()
                db.add(CreditRepayment(credit_customer_id=1,incoming_flow_id=r_id,repayment_date=datetime(2026,10,20),amount=1000))
                db.add(AccountTransfer(warehouse_id=1,from_account_id=1,to_account_id=3,transfer_date=date(2026,10,20),amount=300,currency='THB',request_key='round5-transfer',operator_id=1))
                await db.commit()
            records,summary,sheet=await check_scope(client,'income',{'THB':1000},1,customer_id=1)
            assert records['data'][0]['customer_name']=='测试客户'
            await check_scope(client,'expense',{},0)
            actual('receipt_links',{'declaration_only':0,'receipt_with_two_matches_and_credit_link':1000,'rows':1,'internal_transfer_excluded':300,'customer':records['data'][0]['customer_name']})
    asyncio.run(run())


def test_income_filters_share_all_rows_summary_and_export_with_no_writes(tmp_path,monkeypatch):
    async def run():
        import app.services.cost_funds_debt as dates
        monkeypatch.setattr(dates,'thai_today',lambda:date(2026,10,20))
        async with environment(tmp_path,monkeypatch) as (engine,factory,client):
            async with factory() as db:
                db.add(PaymentAccount(id=3,warehouse_id=1,account_name='人民币账户',account_type='cash',account_number='3',currency='CNY'))
                db.add(Customer(id=2,warehouse_id=2,customer_code='FOREIGN',company_name='SECRET-WH2'))
                db.add_all([IncomeExpenseCategory(id=1,warehouse_id=1,type='income',name='业务收款',category_group='operating'),
                    IncomeExpenseCategory(id=2,warehouse_id=1,type='income',name='其他收入',category_group='other')])
                for uid in (3,4): (await db.get(User,uid)).extra_permissions=['运营收支','其他收支']
                await db.flush()
                db.add_all([IncomeRecord(id=1,warehouse_id=1,account_id=1,category_id=1,customer_id=1,income_date=datetime(2026,10,20),amount=200,currency='THB',remark='MANUAL-200'),
                    IncomeRecord(id=2,warehouse_id=1,account_id=1,category_id=2,income_date=datetime(2026,10,20),amount=40,currency='THB',remark='OTHER-40'),
                    IncomeRecord(id=3,warehouse_id=1,account_id=3,category_id=1,income_date=datetime(2026,10,20),amount=50,currency='CNY',remark='CNY-50'),
                    IncomeRecord(id=4,warehouse_id=1,account_id=1,category_id=1,income_date=datetime(2026,9,30,23,59,59),amount=800,currency='THB',remark='SEPTEMBER-800'),
                    IncomingFlow(id=1,warehouse_id=1,received_date=datetime(2026,10,20),amount=1000,currency='THB',account_id=1,entrant_id=1,payer_name='receipt person'),
                    IncomingFlow(id=2,warehouse_id=2,received_date=datetime(2026,10,20),amount=999,currency='THB',account_id=2,entrant_id=7,payer_name='SECRET-WH2')])
                await db.commit()
            statements=[]
            def trace(conn,cursor,statement,*_): statements.append(statement)
            event.listen(engine.sync_engine,'before_cursor_execute',trace)
            scenarios=[({}, {'THB':1240,'CNY':50},4),({'currency':'THB'},{'THB':1240},3),
                ({'currency':'CNY'},{'CNY':50},1),({'account_id':3},{'CNY':50},1),({'account_id':2},{},0),
                ({'source':'incoming'},{'THB':1000},1),({'source':'manual'},{'THB':240,'CNY':50},3),
                ({'category_id':1},{'THB':200,'CNY':50},2),({'category_group':'other'},{'THB':40},1),
                ({'category_group':'operating'},{'THB':1200,'CNY':50},3),({'search':'MANUAL-200'},{'THB':200},1),
                ({'search':'到账流水'},{'THB':1000},1),({'search':'receipt person'},{'THB':1000},1),
                ({'customer_id':1},{'THB':200},1),({'customer_id':2},{},0),({'source':'unknown'},{},0),
                ({'month':'2026-09','start_date':'2026-09-01','end_date':'2026-09-30'},{'THB':800},1),
                ({'start_date':'2026-10-21'},{},0)]
            try:
                default_list=await client.get('/api/v1/income-expense/income',headers=auth(1))
                default_summary=await client.get('/api/v1/income-expense/monthly-summary',headers=auth(1))
                default_export=await client.get('/api/v1/income-expense/income/export',headers=auth(1))
                assert default_list.status_code==default_summary.status_code==default_export.status_code==200
                assert default_list.json()['totals_by_currency']==default_summary.json()['income_by_currency']=={'THB':1240,'CNY':50}
                assert len(list(load_workbook(BytesIO(default_export.content)).active.values))==5
                history=await read(client,'income',start_date='1900-01-01')
                assert history['totals_by_currency']=={'THB':2040,'CNY':50} and history['total']==5
                for filters,amounts,count in scenarios:
                    listed,_,_=await check_scope(client,'income',amounts,count,**filters)
                    assert 'SECRET-WH2' not in json.dumps(listed)
                for uid in (3,4): await check_scope(client,'income',{'THB':1000},1,uid=uid,source='incoming')
                await check_scope(client,'income',{'THB':999},1,uid=7)
                for uid in (2,5):
                    for path in ('income','expense','monthly-summary','income/export','expense/export'):
                        response=await client.get('/api/v1/income-expense/'+path,headers=auth(uid),params=RANGE)
                        assert response.status_code==403,(uid,path,response.text)
                assert not any(s.lstrip().upper().startswith(('INSERT','UPDATE','DELETE')) for s in statements)
            finally: event.remove(engine.sync_engine,'before_cursor_execute',trace)
            actual('income_filter_contract',{'scenarios':len(scenarios),'unfiltered':{'THB':1240,'CNY':50},'explicit_history':{'THB':2040,'CNY':50},'other_warehouse':999,'read_writes':0,'super_labor':403})
    asyncio.run(run())


def test_expense_business_sources_are_once_only_and_invalid_states_excluded(tmp_path,monkeypatch):
    async def run():
        async with environment(tmp_path,monkeypatch) as (_,factory,client):
            async with factory() as db:
                db.add_all([IncomeExpenseCategory(id=1,warehouse_id=1,type='expense',name='运营费用',category_group='operating'),
                    IncomeExpenseCategory(id=2,warehouse_id=1,type='expense',name='其他费用',category_group='other'),
                    Supplier(id=1,warehouse_id=1,name='supplier'),Employee(id=1,warehouse_id=1,name='worker',hire_date=datetime(2026,10,1))])
                db.add(ExpenseFund(id=1,warehouse_id=1,employee_id=4,receive_date=datetime(2026,10,10),amount=1000,purpose='FUND',remaining_balance=800,currency='THB'))
                db.add(CashAdvance(id=1,warehouse_id=1,payer_name='supervisor',advance_date=date(2026,10,1),amount=1000,purpose='salary',status='confirmed',repaid_amount=300,created_by_name='admin'))
                db.add(PayableBill(id=1,warehouse_id=1,supplier_id=1,bill_number='VOID-BUT-PAID-100',bill_date=datetime(2026,9,1),due_date=datetime(2026,11,1),amount=500,paid_amount=100,legacy_paid_amount=0,status='voided'))
                db.add(Reimbursement(id=1,warehouse_id=1,employee_id=4,submit_date=datetime(2026,10,1),total_amount=80,currency='THB',status='paid',funding_source='company_account'))
                await db.flush()
                for id,amount in [(1,20),(2,40),(3,1000),(4,80),(5,700),(6,900),(7,11),(8,12),(9,13)]:
                    db.add(ExpenseRecord(id=id,warehouse_id=1,account_id=1,category_id=2 if id==2 else 1,supplier_id=1 if id==1 else None,expense_date=datetime(2026,10,20),amount=amount,currency='THB',remark='manual-'+str(id)))
                await db.flush()
                (await db.get(Reimbursement,1)).expense_record_id=4
                db.add(ExpenseFundDisbursement(id=1,warehouse_id=1,fund_id=1,amount=1000,currency='THB',disbursed_date=datetime(2026,10,10),payment_account_id=1,account_name='cash',expense_record_id=3,operator_name='admin'))
                db.add(ExpenseFundItem(fund_id=1,expense_date=datetime(2026,10,20),category='supplies',amount=200,currency='THB',description='spend previously funded money',review_status='approved'))
                db.add(ReimbursementPayment(id=1,warehouse_id=1,reimbursement_id=1,amount=80,currency='THB',funding_source='company_account',payment_account_id=1,previous_status='approved',paid_by_name='admin',paid_at=datetime(2026,10,20,tzinfo=timezone.utc),expense_record_id=4,status='active'))
                db.add(ReimbursementPayment(id=2,warehouse_id=1,reimbursement_id=1,amount=90,currency='THB',funding_source='company_account',payment_account_id=1,previous_status='approved',paid_by_name='admin',paid_at=datetime(2026,10,20,tzinfo=timezone.utc),status='revoked'))
                db.add(CashAdvanceRepayment(id=1,warehouse_id=1,cash_advance_id=1,repayment_date=date(2026,10,20),amount=300,account_id=1,account_name='cash',currency='THB',payment_method='cash',voucher_path='test-only',created_by_name='admin'))
                db.add(CashAdvanceRepayment(id=2,warehouse_id=1,cash_advance_id=1,repayment_date=date(2026,10,20),amount=50,account_id=1,account_name='cash',currency='THB',payment_method='cash',voucher_path='test-only',created_by_name='admin',status='revoked'))
                db.add(PayablePayment(id=1,warehouse_id=1,bill_id=1,account_id=1,amount=100,currency='THB',paid_at=datetime(2026,10,20,tzinfo=timezone.utc),operator_id=1))
                db.add(PayrollRecord(id=1,warehouse_id=1,employee_id=1,period='2026-10',status='confirmed',disbursed=True,net_pay=700,cash_advance_id=1,expense_record_id=5))
                db.add(PayrollRecord(id=2,warehouse_id=1,employee_id=1,period='2026-10',status='confirmed',disbursed=False,net_pay=900,expense_record_id=6))
                for id,state in [(7,'pending'),(8,'rejected'),(9,'approved')]:
                    db.add(EmployeeAdvance(id=id,warehouse_id=1,employee_id=1,amount=id+4,advance_date=date(2026,10,20),period='2026-10',
                        approval_status=state,status='voided' if id==9 else 'unpaid',expense_record_id=id))
                await db.commit()
            # 20 + 40 + 1000 + 80 + 300 + 100; not +200 fund spend, +90/+50 revoked or +700 advance-funded salary.
            listed,summary,sheet=await check_scope(client,'expense',{'THB':1540},6)
            for filters,total,count in [({'source':'expense_fund'},1000,1),({'source':'cash_advance_repayment'},300,1),
                ({'source':'payable'},100,1),({'source':'reimbursement'},80,1),({'source':'manual'},60,2),
                ({'category_group':'operating'},1500,5),({'category_group':'other'},40,1),
                ({'account_id':2},0,0),({'cost_category':'采购'},100,1),({'search':'manual-1'},20,1),({'supplier_id':1},20,1),({'supplier_id':2},0,0)]:
                await check_scope(client,'expense',{'THB':total} if total else {},count,**filters)
            assert sum(r['amount'] for r in listed['data'])==1540
            actual('expense_sources',{'expense':1540,'rows':6,'fund_disbursement':1000,'fund_spend_excluded':200,'reimbursement':80,
                'cash_advance_repayment':300,'advance_funded_salary_excluded':700,'paid_part_of_voided_bill':100,'manual':60,'revoked_excluded':140,
                'pending_rejected_voided_advances_excluded':36})
    asyncio.run(run())
