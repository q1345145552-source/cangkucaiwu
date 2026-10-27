"""Downloaded workbooks and HTTP on a disposable localhost schema, synthetic data only."""
import asyncio, io, json, uuid, zipfile
import pytest
from openpyxl import Workbook,load_workbook
from sqlalchemy import select,func,text
from app.models import Supplier,SupplierProduct,IncomingFlow,PaymentAccount
from app.models.supplier import SupplierCrossBorderPrice
from tests.test_supplier_workflow_round1 import setup,request
from tests.test_incoming_operations import URL
from tests.test_round2_permissions import auth
pytestmark=pytest.mark.skipif(not URL,reason='Explicit isolated localhost PostgreSQL required')
MIME='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

def blob(headers,rows):
    wb=Workbook();wb.active.append(headers)
    for row in rows:wb.active.append(row)
    output=io.BytesIO();wb.save(output);return output.getvalue()

def fill(content,values):
    wb=load_workbook(io.BytesIO(content));ws=wb.active;ws.delete_rows(2,ws.max_row)
    for row in values:ws.append(row)
    out=io.BytesIO();wb.save(out);return out.getvalue()

async def send(c,path,content,data=None,code=200,name='synthetic.xlsx',uid=1,warehouse='1'):
    return await request(c,'POST',path,uid=uid,warehouse=warehouse,code=code,files={'file':(name,content,MIME)},data=data or {})

async def count(f,model):
    async with f() as db:return (await db.execute(select(func.count()).select_from(model))).scalar_one()

async def template(c,path):
    result=await c.get('/api/v1'+path,headers={**auth(1),'X-Warehouse-ID':'1'})
    assert result.status_code==200 and result.headers['content-type'].startswith(MIME)
    return result.content

async def confirm(c,path,content,preview,key=None):
    data={'mode':'confirm','request_key':key or uuid.uuid4().hex,'preview_version':preview['preview_version']}
    if path.endswith('/products'):data['choices']=json.dumps([{'row_number':r['row_number'],'confirmed':True,'decision':'unmapped'} for r in preview['rows']])
    return await send(c,path,content,data),data

@pytest.mark.parametrize('path',['/suppliers/1/import/products','/suppliers/import/products'])
def test_download_products_all_entry_routes_duplicate_policy_and_atomicity(tmp_path,monkeypatch,path):
 async def run():
  async with setup(tmp_path,monkeypatch) as (_,f,c):
   async with f() as db:
    await db.execute(text("SELECT setval(pg_get_serial_sequence('supplier_products','id'),(SELECT max(id) FROM supplier_products))"));await db.commit()
   downloaded=await template(c,'/suppliers/import-template/products')
   example=await send(c,path,downloaded);assert not example['valid'] and '没有可导入数据' in example['errors'][-1]['message']
   content=fill(downloaded,[['本仓供应商','合成袋','10cm',None,5,'个','THB']])
   pre=await send(c,path,content);assert pre['valid'] and pre['imported']==0
   before=await count(f,SupplierProduct)
   result,data=await confirm(c,path,content,pre);assert result['imported']==1
   replay=await send(c,path,content,data);assert replay['id']==result['id'] and replay['replayed']
   assert await count(f,SupplierProduct)==before+1
   # Later explicit user policy: append a new quote, never overwrite the old.
   duplicate=await send(c,path,content);assert duplicate['valid'] and duplicate['rows'][0]['append_quote']
   appended,_=await confirm(c,path,content,duplicate);assert appended['imported']==1 and appended['updated']==0
   assert await count(f,SupplierProduct)==before+2
   mixed=fill(downloaded,[['本仓供应商','有效新袋',None,None,2,'个','THB'],['外仓供应商','错误袋',None,None,2,'个','THB']])
   bad=await send(c,path,mixed)
   if path=='/suppliers/import/products':
    assert bad['valid'] and bad['valid_count']==1 and bad['skipped_count']==1 and bad['skipped_rows'][0]['row_number']==3
   else:
    assert not bad['valid'] and bad['errors'][0]['row_number']==3
    await send(c,path,mixed,{'mode':'confirm','request_key':'bad','preview_version':bad['preview_version'],'choices':'[]'},code=422)
   # Unknown vendor skipping must not relax whole-batch rejection of illegal values.
   invalid=fill(downloaded,[['本仓供应商','有效新袋',None,None,2,'个','THB'],['本仓供应商','错误价格',None,None,'=1+1','个','THB']])
   rejected=await send(c,path,invalid);assert not rejected['valid'] and rejected['errors'][0]['row_number']==3
   await send(c,path,invalid,{'mode':'confirm','request_key':'bad-price','preview_version':rejected['preview_version'],'choices':'[]'},code=422)
   assert await count(f,SupplierProduct)==before+2
   # Product names beginning with 示例: are real data unless the full shipped sample matches.
   named=fill(downloaded,[['本仓供应商','示例:真实产品',None,None,3,'个','THB']]);assert (await send(c,path,named))['valid']
 asyncio.run(run())

@pytest.mark.parametrize('legacy,reordered,origin',[(False,False,'深圳仓'),(False,True,'12345'),(True,False,'深圳仓'),(True,True,'12345')])
def test_logistics_headers_concurrent_lost_response_and_exact_fields(tmp_path,monkeypatch,legacy,reordered,origin):
 async def run():
  async with setup(tmp_path,monkeypatch) as (_,f,c):
   downloaded=await template(c,'/suppliers/import-template/logistics')
   headers=['供应商名称','运输方式','货物类型','发货仓库','单价(元/方)','时效','币种'];values=['本仓供应商','陆运','普货',origin,800,'5-7天','CNY']
   if legacy:headers=headers[1:];values=values[1:]
   if reordered:headers=headers[::-1];values=values[::-1]
   content=blob(headers,[values]) if legacy or reordered else fill(downloaded,[values])
   path='/suppliers/1/import/logistics';pre=await send(c,path,content);assert pre['valid'] and await count(f,SupplierCrossBorderPrice)==0
   data={'mode':'confirm','request_key':'concurrent','preview_version':pre['preview_version']}
   results=await asyncio.gather(*(send(c,path,content,data) for _ in range(4)))
   assert len({r['id'] for r in results})==1 and await count(f,SupplierCrossBorderPrice)==1
   assert (await send(c,path,content,data))['replayed'] # first successful response intentionally ignored
   async with f() as db:
    row=(await db.execute(select(SupplierCrossBorderPrice))).scalar_one()
    assert row.origin_warehouse==origin and row.price_per_cbm==800 and row.currency=='CNY' and row.estimated_days=='5-7天'
 asyncio.run(run())

@pytest.mark.parametrize('kind',['wrong_supplier','bad_amount','bad_currency','duplicate_header','missing_header','empty','example','mixed'])
def test_logistics_validation_all_or_nothing(tmp_path,monkeypatch,kind):
 async def run():
  async with setup(tmp_path,monkeypatch) as (_,f,c):
   downloaded=await template(c,'/suppliers/import-template/logistics');headers=['供应商名称','运输方式','货物类型','发货仓库','单价(元/方)','时效','币种'];row=['本仓供应商','陆运','普货','深圳仓',800,'5天','CNY'];rows=[row]
   if kind=='wrong_supplier':row[0]='外仓供应商'
   if kind=='bad_amount':row[4]='深圳仓'
   if kind=='bad_currency':row[6]='5天'
   if kind=='duplicate_header':headers[-1]='时效'
   if kind=='missing_header':headers[4]='错误金额'
   if kind=='empty':rows=[]
   if kind=='mixed':rows=[row,[*row[:4],'invalid',*row[5:]]]
   content=downloaded if kind=='example' else blob(headers,rows)
   path='/suppliers/1/import/logistics';pre=await send(c,path,content);assert not pre['valid'] and pre['errors']
   await send(c,path,content,{'mode':'confirm','request_key':'invalid','preview_version':pre['preview_version']},code=422)
   assert await count(f,SupplierCrossBorderPrice)==0
 asyncio.run(run())

@pytest.mark.parametrize('kind',['missing_account','foreign_account','currency','date','amount','mixed','no_permission'])
def test_incoming_validation_no_partial_registration(tmp_path,monkeypatch,kind):
 async def run():
  async with setup(tmp_path,monkeypatch) as (_,f,c):
   content=await template(c,'/incoming/template');row=['2026-10-20',200,'THB','合成付款人','现金','隔离测试',1];rows=[row]
   if kind=='missing_account':row[-1]=None
   if kind=='foreign_account':row[-1]=2
   if kind=='currency':row[2]='CNY'
   if kind=='date':row[0]='2026-02-30'
   if kind=='amount':row[1]='unknown'
   if kind=='mixed':rows=[row,[*row[:-1],2]]
   content=fill(content,rows);path='/incoming/import-excel'
   if kind=='no_permission':await send(c,path,content,uid=9,code=403);return
   pre=await send(c,path,content);assert not pre['valid'] and pre['errors']
   await send(c,path,content,{'mode':'confirm','request_key':'invalid','preview_version':pre['preview_version']},code=422)
   assert await count(f,IncomingFlow)==0
 asyncio.run(run())

def test_incoming_template_concurrent_replay_and_legitimate_same_business_values(tmp_path,monkeypatch):
 async def run():
  async with setup(tmp_path,monkeypatch) as (_,f,c):
   content=fill(await template(c,'/incoming/template'),[['2026-10-20',200,'THB','合成付款人','现金','隔离测试',1]])
   path='/incoming/import-excel';pre=await send(c,path,content);assert pre['valid'] and await count(f,IncomingFlow)==0
   data={'mode':'confirm','request_key':'receipt-concurrent','preview_version':pre['preview_version']}
   results=await asyncio.gather(*(send(c,path,content,data) for _ in range(4)))
   assert sum(r['imported'] for r in results if not r.get('replayed'))==1
   assert await count(f,IncomingFlow)==1
   async with f() as db:(await db.get(PaymentAccount,1)).status='inactive';await db.commit()
   assert (await send(c,path,content,data))['replayed']
   async with f() as db:(await db.get(PaymentAccount,1)).status='active';await db.commit()
   result,_=await confirm(c,path,content,pre,'independent-business');assert result['imported']==1
   assert await count(f,IncomingFlow)==2
 asyncio.run(run())

@pytest.mark.parametrize('path',['/suppliers/1/import/products','/suppliers/1/import/logistics','/incoming/import-excel'])
@pytest.mark.parametrize('kind',['zero','corrupt','xls','formula','rows','columns','size','expanded','lying_rows','lying_columns'])
def test_file_protection(tmp_path,monkeypatch,path,kind):
 async def run():
  async with setup(tmp_path,monkeypatch) as (_,f,c):
   name='file.xlsx';content=b''
   if kind=='corrupt':content=b'not xlsx'
   if kind=='xls':name='file.xls';content=blob(['金额'],[[1]])
   if kind=='formula':content=blob(['产品名','单价'],[['公式产品','=1+1']]) if path.endswith('/products') else blob(['金额'],[['=1+1']])
   if kind=='rows':content=blob(['金额'],[[1]]*501)
   if kind=='columns':content=blob(['金额']*21,[[1]*21])
   if kind in ('lying_rows','lying_columns'):
    import re
    original=blob(['金额'],[[1]]*501) if kind=='lying_rows' else blob(['金额']*21,[[1]*21])
    output=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source,zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as target:
     for entry in source.infolist():
      data=source.read(entry.filename)
      if entry.filename=='xl/worksheets/sheet1.xml':data=re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1:A1"',data)
      target.writestr(entry.filename,data)
    content=output.getvalue()
   if kind=='size':content=b'x'*(2*1024*1024+1)
   if kind=='expanded':
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:z.writestr('big.xml',b'x'*(20*1024*1024+1))
    content=out.getvalue()
   if kind=='formula' and path.endswith('/products'):
    result=await send(c,path,content,name=name);assert not result['valid'] and result['errors']
   else:await send(c,path,content,name=name,code=400)
   assert await count(f,IncomingFlow)==0 and await count(f,SupplierCrossBorderPrice)==0 and await count(f,SupplierProduct)==1
 asyncio.run(run())

@pytest.mark.parametrize('kind',['success','ambiguous','foreign'])
def test_global_logistics_current_warehouse_exact_supplier_names(tmp_path,monkeypatch,kind):
 async def run():
  async with setup(tmp_path,monkeypatch) as (_,f,c):
   async with f() as db:
    db.add(Supplier(id=3,warehouse_id=1,name='本仓供应商' if kind=='ambiguous' else '合成供应商乙'));await db.commit()
   rows=[['本仓供应商','陆运','普货','12345',20,'5天','THB'],['外仓供应商' if kind=='foreign' else '合成供应商乙','海运','普货','深圳仓',30,None,'CNY']]
   content=fill(await template(c,'/suppliers/import-template/logistics'),rows);path='/suppliers/import/logistics';pre=await send(c,path,content)
   if kind=='success':
    assert pre['valid'];result,_=await confirm(c,path,content,pre);assert result['imported']==2 and await count(f,SupplierCrossBorderPrice)==2
   else:
    assert not pre['valid'] and pre['errors'];await send(c,path,content,{'mode':'confirm','request_key':'global-invalid','preview_version':pre['preview_version']},code=422);assert await count(f,SupplierCrossBorderPrice)==0
 asyncio.run(run())

@pytest.mark.parametrize('path,template_path',[('/incoming/import-excel','/incoming/template'),('/suppliers/1/import/logistics','/suppliers/import-template/logistics'),('/suppliers/1/import/products','/suppliers/import-template/products')])
@pytest.mark.parametrize('sample',[True,False])
def test_empty_and_only_shipped_example_never_save(tmp_path,monkeypatch,path,template_path,sample):
 async def run():
  async with setup(tmp_path,monkeypatch) as (_,f,c):
   content=await template(c,template_path)
   if not sample:content=fill(content,[])
   result=await send(c,path,content);assert not result['valid'] and not result['rows'] and result['errors']
   assert await count(f,IncomingFlow)==0 and await count(f,SupplierCrossBorderPrice)==0 and await count(f,SupplierProduct)==1
 asyncio.run(run())
