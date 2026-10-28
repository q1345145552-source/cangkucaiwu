"""Downloaded XLSX -> preview -> explicit confirm on disposable localhost schemas."""
import asyncio,io,json,uuid,ast,subprocess
from datetime import datetime
from pathlib import Path
import pytest
from openpyxl import load_workbook,Workbook
from sqlalchemy import select,func,text
from app.models import SupplierProduct,SupplierRoutePrice,Supplier,User,StandardSKU
from app.models.supplier import SupplierCrossBorderPrice
from tests.test_supplier_workflow_round1 import setup
from tests.test_excel_import_repair import send,template,fill,confirm,blob,count
from tests.test_round2_permissions import auth
OUT=Path(__file__).resolve().parents[2]/'docs/supplier-templates-tabs-20261028'
PRODUCT=['品名','规格','单位','单价','币种','阶梯起量','阶梯止量','阶梯单价','生效日期','失效日期']
ROUTE=['起点','终点','运输方式','重量或体积起','重量或体积止','计量单位','单价','币种','时效天数','最低收费','备注']
def evidence(name,value):OUT.joinpath(name+'.json').write_text(json.dumps(value,ensure_ascii=False,indent=2,default=str))
def columns(data):
 wb=load_workbook(io.BytesIO(data));return list(next(wb.active.values))
async def old_template(name):
 root=Path(__file__).resolve().parents[2];source=subprocess.check_output(['git','show','ad9311e:backend/app/api/v1/suppliers.py'],cwd=root,text=True)
 fn=next(n for n in ast.parse(source).body if isinstance(n,ast.AsyncFunctionDef) and n.name==name);fn.decorator_list=[]
 from fastapi import Depends
 from fastapi.responses import StreamingResponse
 ns={'User':User,'Depends':Depends,'get_current_user':lambda:None,'check_staff_permission':lambda _:lambda:None,'io':io,'StreamingResponse':StreamingResponse}
 exec(compile(ast.Module(body=[fn],type_ignores=[]),'<baseline-template>','exec'),ns)
 response=await ns[name]();return b''.join([part async for part in response.body_iterator])
def test_actual_old_and_new_download_columns(tmp_path,monkeypatch):
 async def run():
  async with setup(tmp_path,monkeypatch) as (_,f,c):
   result={}
   for mode,expected,fn in [('products',PRODUCT,'download_products_template'),('logistics',ROUTE,'download_logistics_template')]:
    before=await old_template(fn);single=await template(c,f'/suppliers/import-template/{mode}');batch=await template(c,f'/suppliers/import-template/{mode}?batch=true')
    assert columns(single)==expected and columns(batch)==['供应商名称',*expected]
    result[mode]={'before':columns(before),'after_single':columns(single),'after_batch':columns(batch)}
   evidence('templates-before-after',result)
 asyncio.run(run())
@pytest.mark.parametrize('batch',[False,True])
def test_download_route_template_two_rows_saved_without_legacy_changes(tmp_path,monkeypatch,batch):
 async def run():
  async with setup(tmp_path,monkeypatch) as (_,f,c):
   data=await template(c,f'/suppliers/import-template/logistics?batch={str(batch).lower()}')
   rows=[['深圳仓','曼谷','陆运',0,100,'kg',8,'THB',3,50,'合成按公斤'],['123','春武里','海运',1,20,'方',800,'CNY',7,0,'合成按方']]
   if batch:rows=[['本仓供应商',*row] for row in rows]
   content=fill(data,rows);path='/suppliers/import/logistics' if batch else '/suppliers/1/import/logistics'
   old=await count(f,SupplierCrossBorderPrice);preview=await send(c,path,content);assert preview['valid'],preview
   assert preview['rows'][0]['unit_price']==8 and preview['rows'][1]['origin']=='123' and await count(f,SupplierRoutePrice)==0
   saved,body=await confirm(c,path,content,preview);assert saved['imported']==2
   retry=await send(c,path,content,body);assert retry['replayed'] and await count(f,SupplierRoutePrice)==2 and await count(f,SupplierCrossBorderPrice)==old
   displayed=await c.get('/api/v1/suppliers/1/logistics-prices',headers=auth(1));assert displayed.status_code==200,displayed.text
   assert len(displayed.json()['route_data'])==2 and {r['pricing_unit'] for r in displayed.json()['route_data']}=={'kg','方'}
   evidence('logistics-batch' if batch else 'logistics-single',{'downloaded_columns':columns(data),'preview':preview,'saved':saved,'retry':retry,'display':displayed.json()})
 asyncio.run(run())
@pytest.mark.parametrize('batch',[False,True])
def test_download_product_template_one_row(tmp_path,monkeypatch,batch):
 async def run():
  async with setup(tmp_path,monkeypatch) as (_,f,c):
   data=await template(c,f'/suppliers/import-template/products?batch={str(batch).lower()}')
   values=['合成袋-新模板','38×52cm','个',1.25,'THB',None,None,None,None,None]
   content=fill(data,[[*(['本仓供应商'] if batch else []),*values]])
   path='/suppliers/import/products' if batch else '/suppliers/1/import/products'
   async with f() as db:await db.execute(text("SELECT setval(pg_get_serial_sequence('supplier_products','id'),(SELECT max(id) FROM supplier_products))"));await db.commit()
   before=await count(f,SupplierProduct);preview=await send(c,path,content);assert preview['valid'],preview
   saved,body=await confirm(c,path,content,preview);assert saved['imported']==1 and await count(f,SupplierProduct)==before+1
   assert (await send(c,path,content,body))['replayed'] and await count(f,SupplierProduct)==before+1
   evidence('products-batch' if batch else 'products-single',{'downloaded_columns':columns(data),'preview':preview,'saved':saved})
 asyncio.run(run())
@pytest.mark.parametrize('kind',['columns','unit','range','currency','supplier','amount','formula','days'])
def test_route_invalid_batch_never_partially_saves(tmp_path,monkeypatch,kind):
 async def run():
  async with setup(tmp_path,monkeypatch) as (_,f,c):
   good=['深圳仓','曼谷','陆运',0,100,'kg',8,'THB',3,50,'合成']
   bad=list(good);headers=list(ROUTE)
   if kind=='columns':headers[0]='品名'
   if kind=='unit':bad[5]='吨'
   if kind=='range':bad[4]=0
   if kind=='currency':bad[7]='USD'
   if kind=='supplier':headers=['供应商名称',*headers];good=['本仓供应商',*good];bad=['外仓供应商',*bad]
   if kind=='amount':bad[6]='非法'
   if kind=='formula':bad[6]='=1+1'
   if kind=='days':bad[8]=1.5
   content=blob(headers,[good,bad]);path='/suppliers/1/import/logistics'
   if kind=='formula':await send(c,path,content,code=400);return
   preview=await send(c,path,content);assert not preview['valid'] and preview['errors'],preview
   await send(c,path,content,{'mode':'confirm','request_key':uuid.uuid4().hex,'preview_version':preview['preview_version']},code=422)
   assert await count(f,SupplierRoutePrice)==0
 asyncio.run(run())
def test_reordered_legacy_columns_and_wrong_template_rejection(tmp_path,monkeypatch):
 async def run():
  async with setup(tmp_path,monkeypatch) as (_,f,c):
   path='/suppliers/1/import/logistics'
   old=blob(['币种','时效','单价(元/方)','发货仓库','货物类型','运输方式'],[['CNY','5-7天',800,'123','普货','陆运']])
   p=await send(c,path,old);assert p['valid'] and p['rows'][0]['price_per_cbm']==800,p
   saved,_=await confirm(c,path,old,p);assert saved['imported']==1
   route=blob(list(reversed(ROUTE)),[list(reversed(['深圳仓','曼谷','空运',0,50,'件',5,'THB',1,0,'合成']))]);p=await send(c,path,route);assert p['valid'] and p['rows'][0]['unit_price']==5,p
   product=fill(await template(c,'/suppliers/import-template/products'),[['合成袋','规格','个',5,'THB']])
   p=await send(c,path,product);assert not p['valid'] and '物流模板' in str(p['errors'])
   p=await send(c,'/suppliers/1/import/products',route);assert not p['valid'] and '耗材模板' in str(p['errors'])
 asyncio.run(run())
def test_additive_migration_is_idempotent_and_keeps_old_quotes(tmp_path,monkeypatch):
 async def run():
  from scripts.migrate_supplier_route_prices import migrate
  async with setup(tmp_path,monkeypatch) as (engine,f,c):
   async with f() as db:db.add(SupplierCrossBorderPrice(supplier_id=1,transport_method='陆运',cargo_type='普货',origin_warehouse='深圳仓',price_per_cbm=888,currency='CNY'));await db.commit()
   async with engine.begin() as conn:
    await conn.run_sync(lambda sync:SupplierRoutePrice.__table__.drop(sync));await migrate(conn);await migrate(conn)
   async with f() as db:assert (await db.execute(select(SupplierCrossBorderPrice.price_per_cbm))).scalar_one()==888
   assert await count(f,SupplierRoutePrice)==0
 asyncio.run(run())

def test_optional_tier_dates_reuse_explicit_mapping_and_quote_service(tmp_path,monkeypatch):
 async def run():
  from tests.test_supplier_sku_round5 import fixture,sku,ATTRS
  from tests.test_supplier_workflow_round1 import request
  from app.models import SupplierQuote,SupplierQuoteRevision
  async with fixture(tmp_path,monkeypatch) as (_,f,c,_):
   standard=await sku(c,name='合成阶梯袋')
   data=await template(c,'/suppliers/import-template/products')
   content=fill(data,[['合成阶梯袋','38cm PE','个',5,'THB',1,100,4.5,'2026-10-01','2026-12-31']]);path='/suppliers/1/import/products'
   p=await send(c,path,content);assert p['valid'] and p['rows'][0]['quote_terms']['tiers'][0]['unit_price']=='4.5',p
   before=await count(f,SupplierProduct)
   body={'mode':'confirm','request_key':uuid.uuid4().hex,'preview_version':p['preview_version'],'choices':json.dumps([{'row_number':3,'confirmed':True,'decision':'unmapped'}])}
   # Filled workbook deletes its example; data starts at Excel row 2.
   body['choices']=json.dumps([{'row_number':2,'confirmed':True,'decision':'unmapped'}])
   await send(c,path,content,body,code=400);assert await count(f,SupplierProduct)==before
   candidate=next(v for v in p['rows'][0]['candidates'] if v['id']==standard['id'])
   body['choices']=json.dumps([{'row_number':2,'confirmed':True,'decision':'map','sku_id':candidate['id'],'expected_sku_version':candidate['version'],'reason':'逐项核实合成袋的规格与基本单位','confirmed_attributes':ATTRS}])
   saved=await send(c,path,content,body);assert saved['imported']==1 and saved['data'][0]['quote_id']
   async with f() as db:
    q=(await db.execute(select(SupplierQuote))).scalars().one();rev=(await db.execute(select(SupplierQuoteRevision).where(SupplierQuoteRevision.quote_id==q.id))).scalars().one()
    assert rev.terms['valid_from']=='2026-10-01' and rev.terms['valid_to']=='2026-12-31' and rev.terms['tiers'][0]['unit_price']=='4.5'
    assert rev.terms['fees']['tax_mode']=='unknown'
   replay=await send(c,path,content,body);assert replay['replayed'] and await count(f,SupplierQuote)==1
   evidence('product-tier-dates',{'preview':p,'saved':saved,'replay':replay})
 asyncio.run(run())

def test_existing_pricing_purchase_and_inventory_logic_is_unchanged():
 root=Path(__file__).resolve().parents[2]
 def old(path):return subprocess.check_output(['git','show','ad9311e:'+path],cwd=root,text=True)
 for path in ['backend/app/services/supplier_quotes.py','backend/app/services/inventory.py','backend/app/services/account_balances.py','backend/app/services/supplier_comparison.py']:
  assert (root/path).read_text()==old(path)
 path='backend/app/api/v1/suppliers.py'
 def functions(source):return {n.name:ast.dump(n,include_attributes=False) for n in ast.parse(source).body if isinstance(n,(ast.FunctionDef,ast.AsyncFunctionDef))}
 original=functions(old(path));current=functions((root/path).read_text())
 for name in original:
  if name not in ('download_products_template','download_logistics_template','list_logistics_prices'):assert current[name]==original[name],name

@pytest.mark.parametrize('mode',['products','logistics'])
def test_reordered_example_only_template_has_no_real_rows(tmp_path,monkeypatch,mode):
 async def run():
  async with setup(tmp_path,monkeypatch) as (_,f,c):
   downloaded=await template(c,f'/suppliers/import-template/{mode}')
   wb=load_workbook(io.BytesIO(downloaded));values=list(wb.active.values)
   data=blob(list(values[0])[::-1],[list(values[1])[::-1]])
   result=await send(c,f'/suppliers/1/import/{mode}',data)
   assert not result['valid'] and ('没有实际数据' in str(result) or '没有可导入数据' in str(result)),result
 asyncio.run(run())
