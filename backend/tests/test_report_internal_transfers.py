"""Native PostgreSQL: transfer disclosure never changes company-wallet money."""
import asyncio, json, os
from datetime import date
from io import BytesIO
from pathlib import Path
import pytest
from openpyxl import load_workbook
from sqlalchemy import select,func
from app.models import AccountTransfer, PaymentAccount, UserWarehouse
from tests.test_cash_wallet import cash_environment, balances, receipt, payment, RANGE
from tests.test_incoming_operations import URL
from tests.test_round2_permissions import auth
pytestmark=pytest.mark.skipif(not URL,reason='Requires explicit local PostgreSQL')
ACTUAL={}
PAGES=['/income-expense/monthly-summary','/reports/income-expense','/income-expense/ledger',
 '/reports/cost-funds-debt','/income-expense/operating-dashboard','/dashboard/warehouse-summary','/dashboard/cockpit']
def actual(key,value):
 ACTUAL[key]=value
 if os.environ.get('TRANSFER_REPORT_JSON'):Path(os.environ['TRANSFER_REPORT_JSON']).write_text(json.dumps(ACTUAL,ensure_ascii=False,indent=2))
def transfer(a,b,amount=10000,key='report-main',day='2026-10-20'):
 return dict(from_account_id=a,to_account_id=b,amount=amount,request_key=key,transfer_date=day)
async def setup(f,c,cash,bank):
 async with f() as db:
  for aid,n in [(cash,20000),(bank,3000)]:(await db.get(PaymentAccount,aid)).opening_balance=n
  db.add(UserWarehouse(user_id=1,warehouse_id=2));await db.commit()
 for kind,payload in [('income',receipt(cash)),('expense',payment(cash))]:
  r=await c.post('/api/v1/income-expense/'+kind,headers=auth(1),**({'json':payload} if kind=='income' else {'data':payload}));assert r.status_code==200,r.text
 yuan=[]
 for name,n in [('人民币现金',1000),('人民币银行',500)]:
  r=await c.post('/api/v1/accounts',headers=auth(1),json=dict(account_name=name,account_type='cash',account_number='',opening_balance=n,opening_balance_date='2026-10-01',currency='CNY'))
  assert r.status_code==200,r.text;yuan.append(r.json()['id'])
 return yuan
async def views(c,warehouse=1,month='2026-10'):
 import calendar
 start=date.fromisoformat(month+'-01');params=dict(start_date=str(start),end_date=str(start.replace(day=calendar.monthrange(start.year,start.month)[1])))
 result={}
 for p in PAGES:
  r=await c.get('/api/v1'+p,headers={**auth(1),'X-Warehouse-ID':str(warehouse)},params={'month':month} if p.endswith('cockpit') else params)
  assert r.status_code==200,(p,r.text);result[p]=r.json()
 return result
def stats(p,d):
 if p.endswith('cockpit'):return d['finance']['month']['internal_transfers']
 if p.endswith('warehouse-summary'):return d['totals']['internal_transfers']
 return d['internal_transfers']
def core(p,d):
 if p.endswith('cost-funds-debt'):return {b['currency']:{'cost':b['costs'],'debt':b['debt'],'reconciliation':b['reconciliation'],'closing_balance':b['funds']['total']['closing_balance']} for b in d['currencies']}
 if p.endswith('cockpit'):return {k:d['finance']['month'][k] for k in ('income','expense','net','operating_expense','procurement_expense')}
 if p.endswith('warehouse-summary'):d=d['totals']
 return {k:d[k] for k in ('income_by_currency','expense_by_currency','total_income','total_expense','net')}
def test_nonzero_all_views_balances_retry(tmp_path,monkeypatch):
 async def run():
  async with cash_environment(tmp_path,monkeypatch) as (_,f,c,cash,bank):
   yuan=await setup(f,c,cash,bank);before=await views(c);balance_before=await balances(f,cash,bank,*yuan)
   assert balance_before=={cash:20800,bank:3000,yuan[0]:1000,yuan[1]:500}
   r=await c.post('/api/v1/accounts/transfers',headers=auth(1),json=transfer(cash,bank));assert r.status_code==200,r.text
   thb=await views(c)
   for p,d in thb.items():
    assert core(p,d)==core(p,before[p]),p
    assert stats(p,d)['by_currency']=={'THB':10000},p
    assert stats(p,d)['count_by_currency']=={'THB':1},p
   for _ in range(3):
    retry=await c.post('/api/v1/accounts/transfers',headers=auth(1),json=transfer(cash,bank));assert retry.status_code==200 and retry.json()['id']==r.json()['id']
   r2=await c.post('/api/v1/accounts/transfers',headers=auth(1),json=transfer(*yuan,amount=200,key='report-cny'));assert r2.status_code==200,r2.text
   after=await views(c)
   for p,d in after.items():
    assert core(p,d)==core(p,before[p]),p
    assert stats(p,d)['by_currency']=={'CNY':200,'THB':10000},p
    assert stats(p,d)['total_count']==2
   final=await balances(f,cash,bank,*yuan);assert final=={cash:10800,bank:13000,yuan[0]:800,yuan[1]:700}
   async with f() as db:assert (await db.execute(select(func.count(AccountTransfer.id)))).scalar_one()==2
   actual('nonzero_before_after',dict(before=before,after_thb=thb,after_cny_and_retries=after,balance_before=balance_before,balance_after=final,thb_total_before_after=[23800,23800],cny_total_before_after=[1500,1500],create_responses=[r.json(),r2.json()]))
 asyncio.run(run())
def test_date_warehouse_account_currency_exports(tmp_path,monkeypatch):
 async def run():
  async with cash_environment(tmp_path,monkeypatch) as (_,f,c,cash,bank):
   yuan=await setup(f,c,cash,bank)
   for payload in [transfer(cash,bank),transfer(*yuan,amount=200,key='report-cny')]:assert (await c.post('/api/v1/accounts/transfers',headers=auth(1),json=payload)).status_code==200
   async with f() as db:
    db.add(PaymentAccount(id=99,warehouse_id=2,account_name='外仓银行',account_type='bank',account_number='-',currency='THB',opening_balance=0));await db.flush()
    db.add_all([AccountTransfer(warehouse_id=1,from_account_id=cash,to_account_id=bank,amount=777,currency='THB',transfer_date=date(2026,9,20),request_key='september',operator_id=1),AccountTransfer(warehouse_id=2,from_account_id=2,to_account_id=99,amount=999,currency='THB',transfer_date=date(2026,10,20),request_key='other-wh',operator_id=7)]);await db.commit()
   october=await views(c)
   for items,expected in [(october,{'CNY':200,'THB':10000}),(await views(c,month='2026-09'),{'THB':777}),(await views(c,warehouse=2),{'THB':999})]:
    for p,d in items.items():assert stats(p,d)['by_currency']==expected,(p,d)
   for path in ['/income-expense/ledger','/reports/cost-funds-debt/flows','/income-expense/monthly-summary']:
    r=await c.get('/api/v1'+path,headers=auth(1),params={**RANGE,'account_id':cash,'currency':'THB'});assert r.status_code==200,r.text
    d=r.json();assert d['internal_transfers']['by_currency']=={'THB':10000};assert d['internal_transfers']['scope_label']=='涉及该账户的调拨额'
    assert (await c.get('/api/v1'+path,headers=auth(1),params={**RANGE,'account_id':2})).status_code==403
   exports={}
   for path in ['/reports/income-expense','/income-expense/ledger/export','/income-expense/income/export','/income-expense/expense/export']:
    r=await c.get('/api/v1'+path,headers={**auth(1),'X-Warehouse-ID':'1'},params={**RANGE,'format':'excel'});assert r.status_code==200,(path,r.text)
    book=load_workbook(BytesIO(r.content));rows=list(book['内部转账（不计收支）'].values);assert {r[1]:r[3] for r in rows[1:]}=={'CNY':200,'THB':10000},(path,rows)
    exports[path]=[list(row) for row in rows]
   r=await c.get('/api/v1/income-expense/ledger/export',headers=auth(1),params={**RANGE,'account_id':bank,'currency':'THB','flow_type':'income','source':'transfer'})
   rows=list(load_workbook(BytesIO(r.content))['内部转账（不计收支）'].values);assert {r[1]:r[3] for r in rows[1:]}=={'CNY':0,'THB':10000}
   independent=await c.post('/api/v1/accounts/transfers',headers=auth(1),json=transfer(bank,cash,key='same-amount-new-transfer'))
   assert independent.status_code==200,independent.text
   independent_views=await views(c)
   for p,d in independent_views.items():
    assert stats(p,d)['by_currency']=={'CNY':200,'THB':20000}
    assert stats(p,d)['count_by_currency']=={'CNY':1,'THB':2}
   actual('scope_and_exports',dict(selected_october={p:stats(p,d) for p,d in october.items()},september_thb=777,other_warehouse_october_thb=999,exports=exports,filtered_export=[list(r) for r in rows]))
 asyncio.run(run())
def test_roles_unauthorized_scope_retired(tmp_path,monkeypatch):
 async def run():
  async with cash_environment(tmp_path,monkeypatch) as (_,f,c,cash,bank):
   await setup(f,c,cash,bank);denied=[]
   for p in PAGES:
    for uid in (2,5):
     r=await c.get('/api/v1'+p,headers=auth(uid));assert r.status_code==403,(uid,p,r.text)
    r=await c.get('/api/v1'+p,headers={**auth(1),'X-Warehouse-ID':'999'});assert r.status_code==403,(p,r.text);denied.append(dict(path=p,status=r.status_code,body=r.json()))
    for uid in (1,3,4):
     r=await c.get('/api/v1'+p,headers=auth(uid));assert r.status_code==200,(uid,p,r.text)
   for p in ['/reports/operations','/customers/1/receivables','/credit/shipments']:assert (await c.get('/api/v1'+p,headers=auth(1))).status_code==403
   actual('guards',denied)
 asyncio.run(run())
