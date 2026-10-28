import asyncio,json
from datetime import datetime
from io import BytesIO
import pytest
from openpyxl import load_workbook
from sqlalchemy import select,update
from app.models import User,UserWarehouse,SystemSetting
from app.models.data_change_history import DataChangeHistory
from app.services import permission_migrations as migrations
from tests.test_round2_permissions import auth,set_perms
from tests.test_round3_permissions import hr_environment,snapshot

async def seed_logs(factory):
 async with factory() as db:
  db.add_all([
   DataChangeHistory(id=101,module='recharge',record_id=8,operator_id=1,operator_name='仓库管理员',operation_type='edit',warehouse_id=1,before_data={'amount':10},after_data={'amount':20},created_at=datetime(2026,10,18,9)),
   DataChangeHistory(id=102,module='employee',record_id=9,operator_id=4,operator_name='主管',operation_type='create',warehouse_id=1,after_data={'name':'测试员工'},created_at=datetime(2026,10,18,10)),
   DataChangeHistory(id=201,module='recharge',record_id=8,operator_id=6,operator_name='外仓操作人',operation_type='edit',warehouse_id=2,after_data={'amount':999},created_at=datetime(2026,10,18,11))])
  await db.commit()

async def history_snapshot(factory):
 async with factory() as db:return [tuple(r) for r in (await db.execute(select(DataChangeHistory.__table__).order_by(DataChangeHistory.id))).all()]

@pytest.mark.parametrize('uid,grant',[(1,False),(2,False),(3,False),(3,True),(4,False),(4,True),(5,True)])
def test_same_permission_for_list_operators_export(tmp_path,monkeypatch,uid,grant):
 async def run():
  async with hr_environment(tmp_path,monkeypatch) as (_,f,c):
   await seed_logs(f);before=await history_snapshot(f)
   if grant:await set_perms(f,uid,['修改日志'])
   for path in ('/history','/history/operators','/history/export'):
    r=await c.get('/api/v1'+path,headers=auth(uid))
    allowed=uid in (1,2) or (uid in (3,4) and grant)
    assert r.status_code==(200 if allowed else 403),r.text
    if not allowed:assert r.json()['detail'].startswith('无此操作权限')
    elif path=='/history':assert r.json()['total']==2 and {x['id'] for x in r.json()['data']}=={101,102}
    elif path.endswith('operators'):assert {x['operator_id'] for x in r.json()['data']}=={1,4}
    else:
     rows=list(load_workbook(BytesIO(r.content)).active.values)
     assert {x[0] for x in rows[1:]}=={101,102}
   assert await history_snapshot(f)==before
   if uid in (3,4) and grant:
    await set_perms(f,uid,[])
    for path in ('/history','/history/operators','/history/export'):
     r=await c.get('/api/v1'+path,headers=auth(uid));assert r.status_code==403
 asyncio.run(run())


def test_migration_once_supervisor_only_and_real_read(tmp_path,monkeypatch):
 async def run():
  import scripts.seed as startup
  async with hr_environment(tmp_path,monkeypatch) as (engine,f,c):
   await seed_logs(f)
   await set_perms(f,4,['考勤'])  # Known legacy grants remain eligible; empty grants must stay empty.
   async with f() as db:
    db.add_all([SystemSetting(warehouse_id=1,key=key,value='true') for key in (
     migrations.ROUND2_PERMISSION_MARKER,migrations.ROUND3_PERMISSION_MARKER,migrations.ROUND4_PERMISSION_MARKER,migrations.ROUND5_PERMISSION_MARKER,migrations.ROUND6_PERMISSION_MARKER,migrations.ROUND9_SUPERVISOR_PERMISSION_MARKER,migrations.CASH_ADVANCES_PERMISSION_MARKER)])
    db.add(User(id=7,username='inactive_sup',display_name='停用主管',role='supervisor',warehouse_id=1,password_hash='unused',is_active=False,extra_permissions=['考勤']))
    await db.commit()
    before={u.id:u.extra_permissions for u in (await db.execute(select(User))).scalars()}
   logs=await history_snapshot(f);payroll=await snapshot(f)
   async def seed_noop():pass
   monkeypatch.setattr(startup,'seed',seed_noop);monkeypatch.setattr(startup,'_get_engine',lambda:engine)
   await startup.initialize();await startup.initialize()
   async with f() as db:
    users={u.id:u for u in (await db.execute(select(User))).scalars()}
    for uid in (4,7):assert users[uid].extra_permissions.count('修改日志')==1
    for uid in (1,2,3,5,6):assert users[uid].extra_permissions==before[uid]
    marks=(await db.execute(select(SystemSetting).where(SystemSetting.key==migrations.MODIFICATION_LOGS_PERMISSION_MARKER))).scalars().all();assert len(marks)==1
    print('LOG_MIGRATION',json.dumps({'supervisor':users[4].extra_permissions,'finance':users[3].extra_permissions,'marker':marks[0].key},ensure_ascii=False))
   sup=await c.get('/api/v1/history',headers=auth(4));assert sup.status_code==200 and sup.json()['total']==2
   await set_perms(f,3,['修改日志']);staff=await c.get('/api/v1/history',headers=auth(3));assert staff.status_code==200 and staff.json()['total']==2
   print('LOG_READ_RESULTS',json.dumps({'supervisor':sup.json(),'finance':staff.json()},ensure_ascii=False))
   await set_perms(f,4,[]);await startup.initialize()
   assert (await c.get('/api/v1/history',headers=auth(4))).status_code==403
   assert await history_snapshot(f)==logs and await snapshot(f)==payroll
 asyncio.run(run())


def test_migration_failure_rolls_back(tmp_path,monkeypatch):
 async def run():
  async with hr_environment(tmp_path,monkeypatch) as (engine,f,c):
   async with f() as db:
    db.add(User(id=7,username='synthetic-known-legacy',display_name='SYNTHETIC',role='supervisor',warehouse_id=1,password_hash='unused',extra_permissions=['考勤']))
    await db.commit()
   with pytest.raises(RuntimeError):
    async with engine.begin() as conn:
     await migrations.backfill_modification_logs_permission(conn)
     raise RuntimeError('abort')
   async with f() as db:
    assert (await db.get(User,4)).extra_permissions==[]
    assert (await db.get(User,7)).extra_permissions==['考勤']
    assert not (await db.execute(select(SystemSetting).where(SystemSetting.key==migrations.MODIFICATION_LOGS_PERMISSION_MARKER))).first()
   async with engine.begin() as conn:assert await migrations.backfill_modification_logs_permission(conn)==1
   async with engine.begin() as conn:assert await migrations.backfill_modification_logs_permission(conn)==0
 asyncio.run(run())


def test_warehouse_filter_and_export_exactly_match(tmp_path,monkeypatch):
 async def run():
  async with hr_environment(tmp_path,monkeypatch) as (_,f,c):
   await seed_logs(f);await set_perms(f,4,['修改日志'])
   async with f() as db:
    u=await db.get(User,6);u.role='staff';u.extra_permissions=['修改日志']
    db.add(DataChangeHistory(id=103,module='recharge',record_id=8,operator_id=1,operator_name='=1+1',operation_type='edit',warehouse_id=1,after_data={'amount':30},created_at=datetime(2026,10,17,9)))
    await db.commit()
   for query,ids in [('',{101,102,103}),('?module=recharge&record_id=8&operation_type=edit&operator_id=1&start_date=2026-10-18&end_date=2026-10-18',{101}),('?operator_id=6',set())]:
    r=await c.get('/api/v1/history'+query,headers={**auth(4),'X-Warehouse-ID':'2'});assert r.status_code==200
    x=await c.get('/api/v1/history/export'+query,headers=auth(4));assert x.status_code==200
    assert {v['id'] for v in r.json()['data']}==ids
    assert {v[0] for v in list(load_workbook(BytesIO(x.content)).active.values)[1:]}==ids
   r=await c.get('/api/v1/history',headers=auth(6));assert r.status_code==200 and {x['id'] for x in r.json()['data']}=={201}
   r=await c.get('/api/v1/history/operators',headers=auth(6));assert {x['operator_id'] for x in r.json()['data']}=={6}
   r=await c.get('/api/v1/history/export?end_date=2026-10-17',headers=auth(4));book=load_workbook(BytesIO(r.content));assert book.active['E2'].value=='=1+1' and book.active['E2'].data_type=='s'
 asyncio.run(run())
