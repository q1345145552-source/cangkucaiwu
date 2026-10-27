"""Parse actual generated files, including ZIP metadata and every workbook sheet."""
from pathlib import Path
from io import BytesIO
import json,hashlib,zipfile,sys
from openpyxl import load_workbook
sys.path.append('/Users/xiongliu/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/lib/python3.12/site-packages')
from pypdf import PdfReader
ROOT=Path(__file__).resolve().parents[3];OUT=ROOT/'docs/permissions-round5-20261028';rows=[]
for p in sorted((OUT/'target-07').glob('*.xlsx')):
 wb=load_workbook(p,data_only=False);cells=[str(c.value) for sheet in wb for row in sheet for c in row if c.value is not None]
 assert wb.active.max_row>1
 for text in cells:
  assert not any(x in text for x in ('BANK-SENTINEL','PASSPORT-SENTINEL','LOG-BANK-SENTINEL','SYNTHETIC EMP 8'))
 with zipfile.ZipFile(p) as z:
  for name in z.namelist():
   content=z.read(name)
   assert b'BANK-SENTINEL' not in content and b'SYNTHETIC EMP 8' not in content
 if 'payroll' in p.name:assert 'SYNTHETIC EMP 1' in cells and '100' in cells
 if 'recharge' in p.name:assert '101' in cells and '102' in cells and '108' not in cells
 rows.append({'file':str(p.relative_to(OUT)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'sheets':[{'title':s.title,'rows':s.max_row,'columns':s.max_column,'state':s.sheet_state} for s in wb],'all_sheet_and_zip_metadata_checked':True})
for p in sorted((OUT/'target-07').glob('*.pdf')):
 pdf=PdfReader(BytesIO(p.read_bytes()));text=''.join(page.extract_text() for page in pdf.pages)
 assert 'PO-SENTINEL-1' in text and 'PO-SENTINEL-2' not in text
 assert all(s not in text for s in ('BANK-SENTINEL','PASSPORT-SENTINEL'))
 rows.append({'file':str(p.relative_to(OUT)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'pages':len(pdf.pages),'actual_text':text,'metadata':dict(pdf.metadata or {})})
(OUT/'export-verification.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2));print('parsed and verified',len(rows),'real files')
