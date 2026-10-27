"""Bounded XLSX reader shared by import entry points. Never evaluates formulas."""
import io
import zipfile
from itertools import islice
from fastapi import HTTPException
from openpyxl import load_workbook

async def read_upload(file):
    if not (file.filename or '').lower().endswith('.xlsx'):
        raise HTTPException(400, '仅支持.xlsx；.xls请在Excel中另存为.xlsx，不能只改扩展名')
    content = await file.read(2 * 1024 * 1024 + 1)
    if not content:
        raise HTTPException(400, '文件为空，请选择包含表头和数据的 xlsx 文件')
    if len(content) > 2 * 1024 * 1024:
        raise HTTPException(400, '请选择2MB以内的.xlsx文件')
    return content

def workbook(content, allow_active_formulas=False):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError()
            if '[Content_Types].xml' not in archive.namelist() or 'xl/workbook.xml' not in archive.namelist():
                raise ValueError()
            content_types = archive.read('[Content_Types].xml')
            if b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml' not in content_types or b'macroEnabled' in content_types:
                raise ValueError()
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        for sheet in wb:
            if (sheet.max_row or 0) > 501 or (sheet.max_column or 0) > 20:
                raise ValueError()
            # Uploaded dimension metadata may understate actual rows/columns.
            sheet.reset_dimensions()
            for number, row in enumerate(islice(sheet.iter_rows(), 502), 1):
                if number > 501 or len(row) > 20:
                    raise ValueError()
                if any(c.data_type == 'f' for c in row) and not (allow_active_formulas and sheet == wb.active):
                    raise HTTPException(400, f'{sheet.title}第{number}行包含公式；请粘贴核实后的数值')
        return wb
    except HTTPException:
        if 'wb' in locals(): wb.close()
        raise
    except Exception:
        if 'wb' in locals(): wb.close()
        raise HTTPException(400, '无效或损坏的.xlsx文件，或超过500行、20列、20MB展开限制；未导入')

def table(content, headers, required):
    wb = workbook(content)
    try:
        cells = list(wb.active.iter_rows(values_only=True))
        columns, errors, rows = {}, [], []
        for index, value in enumerate(cells[0] if cells else []):
            if value in (None, ''): continue
            key = headers.get(str(value).strip())
            if key is None: errors.append({'row_number': 1, 'message': f'未知表头「{value}」，请下载模板并保留表头'})
            elif key in columns: errors.append({'row_number': 1, 'message': f'重复表头「{value}」'})
            else: columns[key] = index
        missing = required - columns.keys()
        if missing: errors.append({'row_number': 1, 'message': '缺少表头：' + '、'.join(next(h for h,k in headers.items() if k == key) for key in sorted(missing))})
        for number, values in enumerate(cells[1:], 2):
            if any(v not in (None, '') for v in values):
                rows.append((number, {key: values[i] if i < len(values) else None for key,i in columns.items()}))
        return rows, errors
    finally: wb.close()
