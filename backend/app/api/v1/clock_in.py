from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Query, Request
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, or_, and_
from app.database import get_db
from app.models.user import User
from app.models.clock_in_records import ClockInRecord
from app.core.permissions import get_current_user, get_wh_id, get_wh_ids, Role
from app.core.timezone import thai_now, thai_today, THAI_TZ
from app.core.messages import t, get_request_lang, session_label
from datetime import datetime, date, time, timedelta
from pydantic import BaseModel
from typing import Optional, List
import os, uuid, base64

router = APIRouter()
UPLOAD_DIR = "/app/uploads"

SESSIONS = {
    1: {"label": "早上上班", "time": time(9, 0)},
    2: {"label": "中午休息结束", "time": time(12, 0)},
    3: {"label": "下午上班", "time": time(13, 0)},
    4: {"label": "下午下班", "time": time(18, 0)},
}

def _get_penalty(session: int, clocked_at: datetime) -> dict:
    """Only session 1 tracks late status. Penalty amount computed at payroll time."""
    if session != 1:
        return {"status": "normal", "penalty_amount": 0}
    t = clocked_at.time()
    if t <= time(9, 5):
        return {"status": "normal", "penalty_amount": 0}
    elif t <= time(9, 30):
        return {"status": "late_half", "penalty_amount": 0}  # amount computed at payroll
    else:
        return {"status": "late_one", "penalty_amount": 0}

@router.post("")
async def clock_in(
    request: Request,
    session: int = Form(...),
    photo_base64: str = Form(None),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    lang = get_request_lang(request)
    if current_user.role != Role.WAREHOUSE_LABOR:
        raise HTTPException(403, t("only_labor_can_clock", lang))
    if session not in SESSIONS:
        raise HTTPException(400, t("invalid_session", lang))

    today = thai_today()
    now = thai_now()

    si = SESSIONS[session]
    label = session_label(session, lang)

    # Check duplicate
    existing = (await db.execute(
        select(ClockInRecord).where(
            ClockInRecord.user_id == current_user.id,
            ClockInRecord.clock_date == today,
            ClockInRecord.session == session,
        )
    )).scalar_one_or_none()
    if existing:
        return {"message": t("already_clocked", lang, label=label), "clocked_in_at": existing.clocked_in_at.isoformat(), "duplicate": True}

    # Save photo
    photo_path = None
    if photo_base64:
        # 解码校验（在 try 之外，确保格式错误能正确报错而不是被吞掉）
        try:
            _, data = photo_base64.split(",", 1) if "," in photo_base64 else ("", photo_base64)
            img_bytes = base64.b64decode(data)
        except Exception:
            raise HTTPException(400, t("invalid_photo", lang))
        try:
            wh_id = str(get_wh_id(current_user) or 0)
            today_str = today.isoformat()
            subdir = os.path.join(UPLOAD_DIR, wh_id, today_str, "clockin")
            os.makedirs(subdir, exist_ok=True)
            fname = f"{uuid.uuid4().hex}.jpg"
            fpath = os.path.join(subdir, fname)
            with open(fpath, "wb") as f:
                f.write(img_bytes)
            photo_path = f"uploads/{wh_id}/{today_str}/clockin/{fname}"
        except Exception:
            pass

    wh_id = get_wh_id(current_user)
    penalty = _get_penalty(session, now)
    record = ClockInRecord(
        user_id=current_user.id,
        warehouse_id=wh_id,
        clock_date=today,
        session=session,
        clocked_in_at=now,
        photo_path=photo_path,
        status=penalty["status"],
        penalty_amount=0,  # computed at payroll time
    )
    db.add(record)
    await db.flush()

    if penalty["status"] != "normal":
        msg = t("clock_in_success_late", lang, label=label)
    else:
        msg = t("clock_in_success", lang, label=label)

    return {
        "message": msg,
        "session": session,
        "clocked_in_at": record.clocked_in_at.isoformat(),
        "status": penalty["status"],
        "penalty_amount": 0,
        "photo_path": photo_path,
    }

@router.get("/today")
async def get_today(
    request: Request,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    lang = get_request_lang(request)
    if current_user.role != Role.WAREHOUSE_LABOR:
        raise HTTPException(403, t("no_permission", lang))
    today = thai_today()
    records = (await db.execute(
        select(ClockInRecord).where(
            ClockInRecord.user_id == current_user.id,
            ClockInRecord.clock_date == today,
        ).order_by(ClockInRecord.session)
    )).scalars().all()
    completed = {r.session: {
        "session": r.session, "label": session_label(r.session, lang),
        "clocked_in_at": r.clocked_in_at.isoformat(), "status": r.status,
        "penalty_amount": r.penalty_amount, "photo_path": r.photo_path,
    } for r in records}
    return {
        "today": today.isoformat(),
        "sessions": [{"session": s, "label": session_label(s, lang), "time": str(v["time"])}
                      for s, v in SESSIONS.items()],
        "completed": completed,
    }

@router.get("/records")
async def list_records(
    start_date: str = None,
    end_date: str = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    from app.models.employee import Employee
    if current_user.role == Role.SUPER_ADMIN:
        raise HTTPException(403, "超级管理员请使用各仓库管理员账号操作")
    if current_user.role not in (Role.WAREHOUSE_ADMIN, Role.WAREHOUSE_LABOR, Role.SUPERVISOR):
        raise HTTPException(403, "无权限")

    # Determine warehouse scope
    active_wh = get_wh_id(current_user)

    # Get employees in scope（含已删除，有打卡记录的行仍显示）
    emp_q = select(Employee).where(Employee.status != "resigned")
    if current_user.role == Role.WAREHOUSE_LABOR:
        wh_id = get_wh_id(current_user)
        # 劳工：优先绑定账号，兜底按显示名+仓库匹配（只能看到自己）
        cond = Employee.user_id == current_user.id
        cond = or_(cond, and_(Employee.name == current_user.display_name, Employee.warehouse_id == wh_id) if wh_id else Employee.name == current_user.display_name)
        emp_q = emp_q.where(cond)
    elif active_wh:
        emp_q = emp_q.where(Employee.warehouse_id == active_wh)
    emps = (await db.execute(emp_q.order_by(Employee.name))).scalars().all()

    # Get clock-in records
    query = select(ClockInRecord)
    if current_user.role == Role.WAREHOUSE_LABOR:
        query = query.where(ClockInRecord.user_id == current_user.id)
    elif active_wh:
        query = query.where(ClockInRecord.warehouse_id == active_wh)
    if start_date:
        query = query.where(ClockInRecord.clock_date >= datetime.strptime(start_date, "%Y-%m-%d").date())
    if end_date:
        query = query.where(ClockInRecord.clock_date <= datetime.strptime(end_date, "%Y-%m-%d").date())
    result = await db.execute(query.order_by(ClockInRecord.clock_date.asc(), ClockInRecord.session))
    records = result.scalars().all()

    # Maps（含未绑定账号的兜底匹配：手机号=用户名 / 姓名=显示名）
    from app.services.employee_match import resolve_employee_user_map
    _, user_to_emp = await resolve_employee_user_map(db, emps)

    uid_set = {r.user_id for r in records}
    users_map = {}
    if uid_set:
        us = (await db.execute(select(User).where(User.id.in_(uid_set)))).scalars().all()
        users_map = {u.id: u.display_name for u in us}

    makeup_by_ids = {r.makeup_by for r in records if r.makeup_by}
    makeup_by_names = {}
    if makeup_by_ids:
        us2 = (await db.execute(select(User).where(User.id.in_(makeup_by_ids)))).scalars().all()
        makeup_by_names = {u.id: u.display_name for u in us2}

    return {
        "employees": [{
            "id": e.id, "name": e.name, "position": e.position,
            "user_id": e.user_id, "status": e.status, "photo_path": e.photo_path,
        } for e in emps],
        "records": [{
            "id": r.id, "user_id": r.user_id, "user_name": users_map.get(r.user_id, ""),
            "employee_id": user_to_emp.get(r.user_id),
            "clock_date": r.clock_date.isoformat(), "session": r.session,
            "label": SESSIONS.get(r.session, {}).get("label", ""),
            "clocked_in_at": r.clocked_in_at.isoformat() if r.clocked_in_at else None,
            "status": r.status, "penalty_amount": r.penalty_amount,
            "photo_path": r.photo_path,
            "is_makeup": bool(r.is_makeup),
            "makeup_by": r.makeup_by,
            "makeup_by_name": makeup_by_names.get(r.makeup_by, ""),
            "makeup_at": r.makeup_at.isoformat() if r.makeup_at else None,
            "makeup_reason": r.makeup_reason,
        } for r in records],
    }

class MakeupCreate(BaseModel):
    employee_id: int
    date: str  # YYYY-MM-DD
    sessions: List[int]  # 1-4
    reason: str

@router.post("/makeup")
async def makeup_clock_in(
    req: MakeupCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """管理员/主管补卡：按标准时段时间生成打卡记录，单独标记为补卡，不参与迟到判定。"""
    from app.models.employee import Employee
    from app.services.employee_match import resolve_employee_user_map

    if current_user.role not in (Role.WAREHOUSE_ADMIN, Role.SUPERVISOR):
        raise HTTPException(403, "只有管理员或主管可以补卡")

    wh_id = get_wh_id(current_user)
    if not wh_id:
        raise HTTPException(400, "请先选择仓库")

    try:
        target = datetime.strptime(req.date, "%Y-%m-%d").date()
    except:
        raise HTTPException(400, "日期格式错误")

    today = thai_today()
    if target > today:
        raise HTTPException(400, "不能补未来的日期")
    if (today - target).days > 60:
        raise HTTPException(400, "不能补超过60天以前的记录，请联系管理员")

    reason = (req.reason or "").strip()
    if not reason:
        raise HTTPException(400, "请填写补卡原因")

    sessions: list[int] = []
    for s in req.sessions:
        if s in SESSIONS and s not in sessions:
            sessions.append(s)
    if not sessions:
        raise HTTPException(400, "请选择要补的时段")

    emp = (await db.execute(
        select(Employee).where(Employee.id == req.employee_id, Employee.warehouse_id == wh_id, Employee.is_deleted == False)
    )).scalar_one_or_none()
    if not emp:
        raise HTTPException(404, "员工不存在")

    emp_to_user, _ = await resolve_employee_user_map(db, [emp])
    uid = emp_to_user.get(emp.id)
    if not uid:
        raise HTTPException(400, "该员工没有关联打卡账号，无法补卡")

    created: list[int] = []
    skipped: list[int] = []
    for s in sessions:
        dup = (await db.execute(
            select(ClockInRecord).where(ClockInRecord.user_id == uid, ClockInRecord.clock_date == target, ClockInRecord.session == s)
        )).scalar_one_or_none()
        if dup:
            skipped.append(s)
            continue
        t_std = SESSIONS[s]["time"]
        clocked = datetime.combine(target, t_std, tzinfo=THAI_TZ)
        db.add(ClockInRecord(
            user_id=uid, warehouse_id=wh_id, clock_date=target, session=s,
            clocked_in_at=clocked, status="normal", penalty_amount=0,
            is_makeup=True, makeup_by=current_user.id, makeup_at=thai_now(), makeup_reason=reason,
        ))
        created.append(s)

    await db.flush()
    if skipped:
        return {"message": f"补卡完成：新增 {len(created)} 段，跳过 {len(skipped)} 段（该时段已有记录）", "created": created, "skipped": skipped}
    return {"message": f"补卡完成：新增 {len(created)} 段", "created": created, "skipped": skipped}

@router.get("/records/export")
async def export_records(
    start_date: str = None,
    end_date: str = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """导出当前时间段内的全部打卡记录（一个员工一天一行），生成 Excel。"""
    from app.models.employee import Employee
    from app.models.attendance import LeaveRequest, RestDay, Absence
    from app.services.employee_match import resolve_employee_user_map
    from app.core.timezone import THAI_TZ
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from fastapi.responses import StreamingResponse
    from urllib.parse import quote

    if current_user.role == Role.SUPER_ADMIN:
        raise HTTPException(403, "超级管理员请使用各仓库管理员账号操作")
    if current_user.role not in (Role.WAREHOUSE_ADMIN, Role.SUPERVISOR):
        raise HTTPException(403, "无权限")

    today = thai_today()
    if start_date:
        try:
            range_start = datetime.strptime(start_date, "%Y-%m-%d").date()
        except:
            raise HTTPException(400, "开始日期格式错误 YYYY-MM-DD")
    else:
        range_start = today.replace(day=1)
    if end_date:
        try:
            range_end = datetime.strptime(end_date, "%Y-%m-%d").date()
        except:
            raise HTTPException(400, "结束日期格式错误 YYYY-MM-DD")
    else:
        range_end = today
    if range_start > range_end:
        raise HTTPException(400, "开始日期不能晚于结束日期")
    # 导出只到当天，未来日期无数据
    range_end = min(range_end, today)

    # 员工范围
    active_wh = get_wh_id(current_user)
    emp_q = select(Employee).where(Employee.status != "resigned")
    if active_wh:
        emp_q = emp_q.where(Employee.warehouse_id == active_wh)
    emps = (await db.execute(emp_q.order_by(Employee.name))).scalars().all()
    emp_ids = [e.id for e in emps]

    emp_to_user, user_to_emp = await resolve_employee_user_map(db, emps)
    uid_set = set(emp_to_user.values())
    username_map = {}
    if uid_set:
        us = (await db.execute(select(User).where(User.id.in_(uid_set)))).scalars().all()
        username_map = {u.id: u.username for u in us}

    wh_ids = get_wh_ids(current_user)

    # 打卡记录
    clock_map = {}
    if uid_set:
        clock_records = (await db.execute(
            select(ClockInRecord).where(
                ClockInRecord.warehouse_id.in_(wh_ids),
                ClockInRecord.user_id.in_(uid_set),
                ClockInRecord.clock_date >= range_start,
                ClockInRecord.clock_date <= range_end,
            ).order_by(ClockInRecord.clock_date, ClockInRecord.session)
        )).scalars().all()
        for cr in clock_records:
            eid = user_to_emp.get(cr.user_id)
            if eid is None:
                continue
            clock_map.setdefault((eid, cr.clock_date), {})[cr.session] = cr

    # 请假 / 休息日 / 未到
    leaves = (await db.execute(select(LeaveRequest).where(
        LeaveRequest.warehouse_id.in_(wh_ids), LeaveRequest.employee_id.in_(emp_ids),
        LeaveRequest.leave_date >= range_start, LeaveRequest.leave_date <= range_end,
        LeaveRequest.status == "approved",
    ))).scalars().all()
    leave_set = {(l.employee_id, l.leave_date) for l in leaves}

    rests = (await db.execute(select(RestDay).where(
        RestDay.warehouse_id.in_(wh_ids), RestDay.employee_id.in_(emp_ids),
        RestDay.rest_date >= range_start, RestDay.rest_date <= range_end,
    ))).scalars().all()
    rest_set = {(r.employee_id, r.rest_date) for r in rests}

    absences = (await db.execute(select(Absence).where(
        Absence.warehouse_id.in_(wh_ids), Absence.employee_id.in_(emp_ids),
        Absence.absence_date >= range_start, Absence.absence_date <= range_end,
    ))).scalars().all()
    absence_set = {(a.employee_id, a.absence_date) for a in absences}

    def fmt_time(dt):
        return dt.astimezone(THAI_TZ).strftime("%H:%M") if dt else ""

    rows = []
    for e in emps:
        total_days = (range_end - range_start).days + 1
        for i in range(total_days):
            dt = range_start + timedelta(days=i)
            sessions = clock_map.get((e.id, dt), {})
            times = [fmt_time(sessions.get(s).clocked_in_at) if s in sessions else "" for s in (1, 2, 3, 4)]
            late = any(sessions.get(s) and sessions.get(s).status in ("late_half", "late_one") for s in (1, 2, 3, 4))
            if (e.id, dt) in leave_set:
                status = "请假"
            elif (e.id, dt) in rest_set:
                status = "休息日"
            elif (e.id, dt) in absence_set:
                status = "缺勤"
            elif sessions:
                status = "迟到" if late else "正常"
            else:
                status = "缺勤"  # 未打卡
            rows.append([
                e.name,
                username_map.get(emp_to_user.get(e.id), ""),
                dt.isoformat(),
                times[0], times[1], times[2], times[3],
                "是" if late else "否",
                status,
            ])

    wb = Workbook()
    ws = wb.active
    ws.title = "打卡记录"
    headers = ["员工姓名", "工号", "日期", "第一次打卡时间", "第二次打卡时间", "第三次打卡时间", "第四次打卡时间", "是否迟到", "出勤状态"]
    ws.append(headers)
    hf = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    hfont = Font(bold=True, color="FFFFFF")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = hf
        cell.font = hfont
    for row in rows:
        ws.append(row)
    # 列宽
    widths = [12, 12, 12, 14, 14, 14, 14, 10, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"打卡记录_{range_start.isoformat()}_{range_end.isoformat()}.xlsx"
    encoded = quote(filename)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
    )

@router.get("/photos")
async def get_photos(
    employee_id: int = Query(...),
    date: str = Query(...),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Admin: get all clock-in photos for a specific employee on a specific date"""
    try:
        if current_user.role == Role.SUPER_ADMIN:
            raise HTTPException(403, "超级管理员请使用各仓库管理员账号操作")
        if current_user.role not in (Role.WAREHOUSE_ADMIN, Role.SUPERVISOR, Role.WAREHOUSE_LABOR):
            raise HTTPException(403, "无权限查看打卡照片")

        from app.models.employee import Employee

        # Try employee_id first, then fallback to user_id lookup
        emp = (await db.execute(
            select(Employee).where(Employee.id == employee_id)
        )).scalar_one_or_none()
        if not emp:
            # Fallback: perhaps employee_id is actually user_id
            emp = (await db.execute(
                select(Employee).where(Employee.user_id == employee_id)
            )).scalar_one_or_none()
        if not emp:
            # Last resort: try matching by user's display_name
            u = (await db.execute(
                select(User).where(User.id == employee_id)
            )).scalar_one_or_none()
            if u:
                emp = (await db.execute(
                    select(Employee).where(Employee.name == u.display_name)
                )).scalar_one_or_none()
        if not emp:
            raise HTTPException(404, "员工不存在")

        wh_ids = get_wh_ids(current_user)
        if emp.warehouse_id not in wh_ids:
            raise HTTPException(403, "无权查看该员工")

        # Get user_id from employee link（含兜底匹配：手机号=用户名 / 姓名=显示名）
        uid = emp.user_id
        if not uid:
            from app.services.employee_match import resolve_employee_user_map
            emp_to_user, _ = await resolve_employee_user_map(db, [emp])
            uid = emp_to_user.get(emp.id)

        # 劳工只能查看自己的打卡照片
        if current_user.role == Role.WAREHOUSE_LABOR and (not uid or uid != current_user.id):
            raise HTTPException(403, "只能查看自己的打卡照片")

        try:
            target_date = datetime.strptime(date, "%Y-%m-%d").date()
        except:
            raise HTTPException(400, "日期格式错误 YYYY-MM-DD")

        records = []
        if uid:
            records = (await db.execute(
                select(ClockInRecord).where(
                    ClockInRecord.user_id == uid,
                    ClockInRecord.clock_date == target_date,
                ).order_by(ClockInRecord.session)
            )).scalars().all()

        sessions_data = {}
        for r in records:
            sessions_data[r.session] = {
                "session": r.session,
                "label": SESSIONS.get(r.session, {}).get("label", ""),
                "clocked_in_at": r.clocked_in_at.isoformat() if r.clocked_in_at else None,
                "status": r.status,
                "penalty_amount": r.penalty_amount,
                "photo_path": r.photo_path,
            }

        all_sessions = []
        for s in [1, 2, 3, 4]:
            if s in sessions_data:
                all_sessions.append(sessions_data[s])
            else:
                all_sessions.append({
                    "session": s,
                    "label": SESSIONS.get(s, {}).get("label", ""),
                    "clocked_in_at": None,
                    "status": "missing",
                    "penalty_amount": 0,
                    "photo_path": None,
                })

        return {
            "employee_id": employee_id,
            "employee_name": emp.name,
            "date": date,
            "sessions": all_sessions,
        }
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(500, f"打卡照片查询失败: {str(e)}")
